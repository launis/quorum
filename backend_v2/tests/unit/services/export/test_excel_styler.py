"""Unit tests for corporate Excel styling routines."""

from __future__ import annotations

import openpyxl

from backend_v2.services.export import ExportService
from backend_v2.services.export.excel_styler import apply_corporate_excel_styling


def _create_sample_workbook() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Column 1", "Column 2 Header Text", "Col 3"])
    ws.append(["Short data", "A very long line of data to test column width expansion", "123"])
    ws.append(["Row 2 data", "Another data point", "456"])
    ws.append(["Row 3 data", "Multi-line\ndata cell\nfor testing", "789"])
    return wb


def test_excel_styler_direct_import() -> None:
    """Verify direct import and invocation of apply_corporate_excel_styling."""
    wb = _create_sample_workbook()
    apply_corporate_excel_styling(wb, ["Summary"])
    ws = wb["Summary"]

    # Header styling
    header_cell = ws.cell(row=1, column=1)
    assert header_cell.fill.start_color.rgb == "001F4E78"
    assert header_cell.font.bold is True
    assert header_cell.font.color.rgb == "00FFFFFF"


def test_apply_corporate_excel_styling_existing_parity() -> None:
    """Verify header fill, zebra fill, borders, and column auto-width calculation on openpyxl workbook."""
    wb_direct = _create_sample_workbook()
    wb_proxy = _create_sample_workbook()

    apply_corporate_excel_styling(wb_direct, ["Summary", "NonExistentSheet"])
    ExportService._apply_corporate_excel_styling(wb_proxy, ["Summary", "NonExistentSheet"])

    ws_direct = wb_direct["Summary"]
    ws_proxy = wb_proxy["Summary"]

    assert ws_direct.freeze_panes == "A2"
    assert ws_proxy.freeze_panes == "A2"
    assert ws_direct.row_dimensions[1].height == 28

    # Row 2 (even row index) should have zebra fill "F9FAFB"
    r2_direct = ws_direct.cell(row=2, column=1)
    r2_proxy = ws_proxy.cell(row=2, column=1)
    assert r2_direct.fill.start_color.rgb == "00F9FAFB"
    assert r2_direct.fill.start_color.rgb == r2_proxy.fill.start_color.rgb

    # Row 3 (odd row index) should have white fill "FFFFFF"
    r3_direct = ws_direct.cell(row=3, column=1)
    r3_proxy = ws_proxy.cell(row=3, column=1)
    assert r3_direct.fill.start_color.rgb == "00FFFFFF"
    assert r3_direct.fill.start_color.rgb == r3_proxy.fill.start_color.rgb

    # Check thin border
    assert r2_direct.border.top.border_style == "thin"
    assert r2_direct.border.top.color.rgb == "00D9D9D9"

    # Check column width calculation
    col2_letter = "B"
    assert ws_direct.column_dimensions[col2_letter].width == ws_proxy.column_dimensions[col2_letter].width
    assert ws_direct.column_dimensions[col2_letter].width > 12


def test_apply_corporate_excel_styling_empty_sheet_list() -> None:
    """Verify applying styling with an empty sheet list is a no-op."""
    wb = _create_sample_workbook()
    apply_corporate_excel_styling(wb, [])
    ws = wb["Summary"]
    # Header should not be modified
    assert ws.cell(row=1, column=1).fill.fill_type is None


def test_apply_corporate_excel_styling_only_nonexistent_sheets() -> None:
    """Verify passing only nonexistent sheet names leaves workbook untouched."""
    wb = _create_sample_workbook()
    apply_corporate_excel_styling(wb, ["NoSuchSheet1", "NoSuchSheet2"])
    ws = wb["Summary"]
    assert ws.cell(row=1, column=1).fill.fill_type is None


def test_apply_corporate_excel_styling_max_width_cap() -> None:
    """Verify column width auto-calculation caps at 50 for very wide cells."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Wide"
    ws.append(["ShortHeader"])
    ws.append(["A" * 120])  # Very long cell
    apply_corporate_excel_styling(wb, ["Wide"])
    assert ws.column_dimensions["A"].width == 50


def test_apply_corporate_excel_styling_empty_worksheet() -> None:
    """Verify styling an empty worksheet with None values computes minimum width 12."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Empty"
    ws.cell(row=1, column=1).value = None
    apply_corporate_excel_styling(wb, ["Empty"])
    assert ws.column_dimensions["A"].width == 12
