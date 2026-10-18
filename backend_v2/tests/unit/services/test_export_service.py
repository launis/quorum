"""Unit tests for ExportService covering Excel and flat CSV exports."""

from __future__ import annotations

import csv
import io
from unittest.mock import patch

import pandas as pd
import pytest
from pydantic import ValidationError

from backend_v2.exceptions import AppException, ErrorCodes
from backend_v2.models.core_base import I18nText
from backend_v2.models.domain.execution import ExecutionRecord, ExecutionStepState
from backend_v2.models.domain.matrix import MatrixClaim, MatrixScale, TDAAssertion
from backend_v2.models.domain.prompt_blocks import (
    MatrixPromptBlock,
    SystemRulePromptBlock,
)
from backend_v2.models.dtos.atom_evaluation import ReasoningStepDTO
from backend_v2.models.dtos.atom_result import AtomResultDTO, ErrorDetailsDTO, HydratedAtomDTO
from backend_v2.models.dtos.causal_discovery import (
    AntiFluffAuditDTO,
    CausalEdgeDTO,
    CausalGraphPayloadDTO,
    CausalNodeDTO,
    CausalRootCauseDiagnosisDTO,
    CausalTdaFusionResultDTO,
    FairScoringBreakdownDTO,
    PrescriptiveRemediationDTO,
)
from backend_v2.models.dtos.export import (
    AtomScaleMetadataDTO,
    ExportDenormalizedFlatRowDTO,
    ExportForensicAtomDTO,
)
from backend_v2.models.dtos.matrix_scorecard import MatrixScorecardRowDTO, ScorecardAtomDTO
from backend_v2.models.dtos.report_data import ReportDataDTO
from backend_v2.models.dtos.step_output import StepOutputDTO
from backend_v2.models.enums import (
    ExecutionStatus,
    LaxSDUIComponentType,
    ReportAtomColumn,
    ReportSheetKey,
    VisualIntent,
)
from backend_v2.models.state import TraceEvent
from backend_v2.models.view.sdui import ParagraphBlock, SduiRadarChartBlock
from backend_v2.services.export_service import ExportService
from backend_v2.services.localization import LocalizationService, ReportHeaderResolver
from backend_v2.tests.fakes.in_memory_repositories import (
    InMemoryPromptBlockRepository,
)


def _build_sample_execution(
    status: ExecutionStatus = ExecutionStatus.PASSED, has_atoms: bool = True
) -> ExecutionRecord:
    step_states: dict[str, ExecutionStepState] = {}
    if has_atoms:
        step_states["stp_0123456789abcdef"] = ExecutionStepState(
            id="stp_0123456789abcdef",
            label="Matrix Step Label",
            status=ExecutionStatus.PASSED,
        )

    return ExecutionRecord(
        id="exe_0123456789abcdef",
        workflow_id="wor_0123456789abcdef",
        organization_id="org_0123456789abcdef",
        target_locale="fi",
        status=status,
        step_states=step_states,
    )


def _build_sample_report_dto(has_atoms: bool = True) -> ReportDataDTO:
    axis = MatrixScorecardRowDTO(
        block_id="blk_0123456789abcdef",
        name="Axis 1",
        label_i18n=I18nText(translations={"fi": "Akseli 1", "en": "Axis 1"}),
        score=4.5,
        scale_max=5.0,
        row_explanation="Explanation for row 1",
        is_evaluative=True,
    )
    chart = SduiRadarChartBlock(axes=[axis])
    results: list[AtomResultDTO] = []
    hydrated_references: dict[str, HydratedAtomDTO] = {}
    if has_atoms:
        results = [
            AtomResultDTO(
                tda_id="tda_0123456789abcdef",
                matrix_id="blk_0123456789abcdef",
                status=ExecutionStatus.PASSED,
                source_quotes="Verbatim quote from source document",
                evaluation_reasoning="Reasoning with several words here",
            )
        ]
        hydrated_references = {
            "tda_0123456789abcdef": HydratedAtomDTO(
                sdui_component=LaxSDUIComponentType.BOOLEAN_CARD,
                resolved_claim="Claim Label",
            )
        }
    return ReportDataDTO(
        workflow_id="wor_0123456789abcdef",
        execution_id="exe_0123456789abcdef",
        profile_id="prof_0123456789abcdef",
        global_score=4.5,
        has_warning=False,
        inner_sdui_blocks=[chart],
        results=results,
        hydrated_references=hydrated_references,
    )


def _build_sample_matrix_rows(n: int = 1) -> list[MatrixScorecardRowDTO]:
    """Helper to construct valid MatrixScorecardRowDTO test instances."""
    return [
        MatrixScorecardRowDTO(
            block_id=f"blk_matrix_{i:04d}",
            name=f"Matrix {i}",
            label_i18n=I18nText(translations={"fi": f"Matriisi {i}", "en": f"Matrix {i}"}),
            score=4.0,
            scale_max=5.0,
            row_explanation=f"Row explanation {i}",
            is_evaluative=True,
        )
        for i in range(1, n + 1)
    ]


@pytest.mark.asyncio
async def test_export_excel_success_fi() -> None:
    service = ExportService()
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED)
    report_dto = _build_sample_report_dto()

    lbl = I18nText(translations={"fi": "L", "en": "L"})
    desc = I18nText(translations={"fi": "D", "en": "D"})
    scale = MatrixScale(score=1, ai_label="L1", claims=[])
    matrix_block = MatrixPromptBlock(
        id="blk_0123456789abcdef",
        slug="m_slug",
        label=lbl,
        description=desc,
        scales=[scale],
        ai_description="Operational rule",
    )
    payload = await service.export_excel(
        execution=exec_record,
        report_dto=report_dto,
        matrices=_build_sample_matrix_rows(1),
        locale="fi",
        components=[matrix_block],
    )

    assert payload.filename == "execution_export_exe_0123456789abcdef.xlsx"
    assert len(payload.content_bytes) > 0
    assert payload.content_bytes.startswith(b"PK")


@pytest.mark.asyncio
async def test_export_excel_success_en_with_prompt_block_repo() -> None:
    prompt_block_repo = InMemoryPromptBlockRepository()
    lbl = I18nText(translations={"fi": "L", "en": "L"})
    desc = I18nText(translations={"fi": "D", "en": "D"})
    scale = MatrixScale(score=1, ai_label="L1", claims=[])
    matrix_block = MatrixPromptBlock(
        id="blk_0123456789abcdef",
        slug="m_slug",
        label=lbl,
        description=desc,
        scales=[scale],
        ai_description="English operational rule",
    )
    await prompt_block_repo.create_prompt_block(matrix_block)

    service = ExportService(prompt_block_repo=prompt_block_repo)
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED)
    report_dto = _build_sample_report_dto()

    payload = await service.export_excel(
        execution=exec_record,
        report_dto=report_dto,
        matrices=_build_sample_matrix_rows(1),
        locale="en",
    )

    assert payload.filename == "execution_export_exe_0123456789abcdef.xlsx"
    assert len(payload.content_bytes) > 0
    assert prompt_block_repo._call_counts["get_all_prompt_blocks"] == 1


@pytest.mark.asyncio
async def test_export_excel_fails_non_passed_status() -> None:
    service = ExportService()
    exec_record = _build_sample_execution(status=ExecutionStatus.FAILED)

    with pytest.raises(AppException) as exc_info:
        await service.export_excel(execution=exec_record, report_dto=None, matrices=_build_sample_matrix_rows(1))

    assert exc_info.value.status_code == 400
    assert exc_info.value.details["error_code"] == ErrorCodes.VALIDATION_FAILED.value
    assert "Execution must be in PASSED state" in exc_info.value.message


@pytest.mark.asyncio
async def test_export_excel_fails_no_scoreable_atoms() -> None:
    service = ExportService()
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED, has_atoms=False)

    with pytest.raises(AppException) as exc_info:
        await service.export_excel(execution=exec_record, report_dto=None, matrices=_build_sample_matrix_rows(1))

    assert exc_info.value.status_code == 400
    assert exc_info.value.details["error_code"] == ErrorCodes.VALIDATION_FAILED.value
    assert "Execution has no scoreable atoms" in exc_info.value.message


@pytest.mark.asyncio
async def test_export_excel_fails_when_report_dto_has_no_atoms() -> None:
    service = ExportService()
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED, has_atoms=False)
    report_dto = _build_sample_report_dto(has_atoms=False)

    with pytest.raises(AppException) as exc_info:
        await service.export_excel(execution=exec_record, report_dto=report_dto, matrices=_build_sample_matrix_rows(1))

    assert exc_info.value.status_code == 400
    assert exc_info.value.details["error_code"] == ErrorCodes.VALIDATION_FAILED.value
    assert "Execution has no scoreable atoms" in exc_info.value.message


@pytest.mark.asyncio
async def test_export_excel_writer_error() -> None:
    service = ExportService()
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED)
    report_dto = _build_sample_report_dto()
    matching_matrix = MatrixScorecardRowDTO(
        block_id="blk_0123456789abcdef",
        name="Axis 1",
        label_i18n=I18nText(translations={"fi": "Akseli 1", "en": "Axis 1"}),
        score=4.5,
        scale_max=5.0,
        row_explanation="Explanation for row 1",
        is_evaluative=True,
    )

    with patch("pandas.ExcelWriter", side_effect=RuntimeError("Disk failure")):
        with pytest.raises(AppException) as exc_info:
            await service.export_excel(execution=exec_record, report_dto=report_dto, matrices=[matching_matrix])

    assert exc_info.value.status_code == 500
    assert exc_info.value.details["error_code"] == ErrorCodes.INTERNAL_SERVER_ERROR.value


@pytest.mark.asyncio
async def test_export_flat_csv_success() -> None:
    service = ExportService()
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED)
    report_dto = _build_sample_report_dto()
    matching_matrix = MatrixScorecardRowDTO(
        block_id="blk_0123456789abcdef",
        name="Axis 1",
        label_i18n=I18nText(translations={"fi": "Akseli 1", "en": "Axis 1"}),
        score=4.5,
        scale_max=5.0,
        row_explanation="Explanation for row 1",
        is_evaluative=True,
    )
    matrices = [matching_matrix]

    payload = await service.export_flat_csv(
        execution=exec_record,
        report_dto=report_dto,
        matrices=matrices,
        locale="fi",
        execution_id="exe_custom_999",
    )

    assert payload.filename == "execution_export_exe_custom_999.csv"
    csv_text = payload.content_bytes.decode("utf-8-sig")
    reader = list(csv.reader(io.StringIO(csv_text)))
    assert len(reader[0]) == 22
    assert len(reader) - 1 == len(report_dto.results)
    assert "Logiikkamatriisi" in reader[0]
    assert "Arviointikriteeri" in reader[0]


@pytest.mark.asyncio
async def test_export_excel_with_report_dto_results_atoms() -> None:
    """Regression test proving failure when atoms are passed in report_dto.results.

    During real runtime DAG execution, ExecutionRecord.step_states has empty scorecard_atoms={}.
    The evaluated atoms exist exclusively as AtomResultDTOs inside report_dto.results and
    hydrated_references. ExportService must extract raw data rows from report_dto.results
    instead of failing with 'Execution has no scoreable atoms'.
    """
    from backend_v2.models.dtos.atom_result import AtomResultDTO, HydratedAtomDTO
    from backend_v2.models.enums import ExecutionStatus

    service = ExportService()
    # Real execution has empty scorecard_atoms in step_states
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED, has_atoms=False)

    atom_result = AtomResultDTO(
        tda_id="tda_0123456789abcdef",
        matrix_id="blk_0123456789abcdef",
        status=ExecutionStatus.PASSED,
        source_quotes="Verbatim quote from source document",
        evaluation_reasoning="Sound reasoning based on evidence.",
    )
    hydrated_ref = HydratedAtomDTO(
        sdui_component="boolean_card",
        resolved_claim="The organization follows clear strategy guidelines.",
    )

    axis = MatrixScorecardRowDTO(
        block_id="blk_0123456789abcdef",
        name="Strategy Matrix",
        label_i18n=I18nText(translations={"fi": "Strategiamatriisi", "en": "Strategy Matrix"}),
        score=4.0,
        scale_max=5.0,
        row_explanation="Solid strategic alignment.",
        is_evaluative=True,
    )
    chart = SduiRadarChartBlock(axes=[axis])

    report_dto = ReportDataDTO(
        workflow_id="wor_0123456789abcdef",
        execution_id="exe_0123456789abcdef",
        profile_id="prof_0123456789abcdef",
        global_score=4.0,
        has_warning=False,
        inner_sdui_blocks=[chart],
        results=[atom_result],
        hydrated_references={"tda_0123456789abcdef": hydrated_ref},
    )

    payload = await service.export_excel(
        execution=exec_record,
        report_dto=report_dto,
        matrices=[axis],
        locale="fi",
    )

    assert payload.filename == "execution_export_exe_0123456789abcdef.xlsx"
    assert len(payload.content_bytes) > 0


@pytest.mark.asyncio
async def test_export_excel_with_all_extensions_and_blocks() -> None:
    """Test Excel export with non-matrix blocks, None labels, all atom extensions, and custom ID."""
    service = ExportService()
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED, has_atoms=False)

    axis_custom = MatrixScorecardRowDTO(
        block_id="blk_aaaaaaaaaaaaaaaa",
        name="Custom Axis",
        label_i18n=I18nText(translations={"en": "Custom Axis"}),
        score=3.0,
        scale_max=5.0,
        row_explanation="Axis with custom label",
        is_evaluative=True,
    )
    chart = SduiRadarChartBlock(axes=[axis_custom])
    text_block = ParagraphBlock(text="Informational text")

    atom_with_extensions = AtomResultDTO(
        tda_id="tda_bbbbbbbbbbbbbbbb",
        matrix_id="blk_aaaaaaaaaaaaaaaa",
        status=ExecutionStatus.FAILED,
        source_quotes=None,
        evaluation_reasoning="Failed check due to missing policy.",
        extensions={
            "internalized_rule": "Strict rule",
            "confidence": "0.85",
            "source_id": "src_custom",
            "falsification": "Refuted claim",
        },
    )
    atom_system_error = AtomResultDTO(
        tda_id="tda_cccccccccccccccc",
        matrix_id="blk_aaaaaaaaaaaaaaaa",
        status=ExecutionStatus.SYSTEM_ERROR,
        source_quotes=None,
        evaluation_reasoning=None,
        error_details=ErrorDetailsDTO(error_code="SYSTEM_ERROR", message="System execution failure."),
    )
    hydrated_ref_with_quote = HydratedAtomDTO(
        sdui_component=LaxSDUIComponentType.BOOLEAN_CARD,
        resolved_claim="Resolved Claim with Quote",
        source_quotes="Quote from hydrated reference",
    )

    component_block = SystemRulePromptBlock(
        id="blk_dddddddddddddddd",
        slug="comp_slug",
        label=I18nText(translations={"fi": "Komponentti", "en": "Component"}),
        description=I18nText(translations={"fi": "D", "en": "D"}),
        instruction_text="Component rule instruction",
    )
    tda_matched_block = SystemRulePromptBlock(
        id="tda_eeeeeeeeeeeeeeee",
        slug="tda_slug",
        label=I18nText(translations={"fi": "TDA Block", "en": "TDA Block"}),
        description=I18nText(translations={"fi": "D", "en": "D"}),
        instruction_text="TDA matched instruction",
    )
    unmatched_block = SystemRulePromptBlock(
        id="blk_0000000000000000",
        slug="unmatched_slug",
        label=I18nText(translations={"fi": "Tuntematon", "en": "Unknown"}),
        description=I18nText(translations={"fi": "D", "en": "D"}),
        instruction_text="Unmatched instruction",
    )

    atom_with_comp_matrix = AtomResultDTO(
        tda_id="tda_eeeeeeeeeeeeeeee",
        matrix_id="blk_dddddddddddddddd",
        status=ExecutionStatus.PASSED,
        source_quotes="Explicit quote",
        evaluation_reasoning="Evaluation passed smoothly.",
    )
    atom_unmatched = AtomResultDTO(
        tda_id="tda_ffffffffffffffff",
        matrix_id="blk_0000000000000000",
        status=ExecutionStatus.PASSED,
        source_quotes="Unknown quote",
        evaluation_reasoning="Evaluation for unknown block.",
    )

    report_dto = ReportDataDTO(
        workflow_id="wor_0123456789abcdef",
        execution_id="exe_0123456789abcdef",
        profile_id="prof_0123456789abcdef",
        global_score=3.5,
        has_warning=False,
        inner_sdui_blocks=[text_block, chart],
        results=[atom_with_extensions, atom_system_error, atom_with_comp_matrix, atom_unmatched],
        hydrated_references={
            "tda_bbbbbbbbbbbbbbbb": hydrated_ref_with_quote,
            "tda_cccccccccccccccc": HydratedAtomDTO(
                sdui_component=LaxSDUIComponentType.BOOLEAN_CARD,
                resolved_claim="Claim for system error",
            ),
            "tda_eeeeeeeeeeeeeeee": HydratedAtomDTO(
                sdui_component=LaxSDUIComponentType.BOOLEAN_CARD,
                resolved_claim="Claim for component matrix",
            ),
            "tda_ffffffffffffffff": HydratedAtomDTO(
                sdui_component=LaxSDUIComponentType.BOOLEAN_CARD,
                resolved_claim="Claim for unknown block",
            ),
        },
    )

    payload = await service.export_excel(
        execution=exec_record,
        report_dto=report_dto,
        matrices=[axis_custom],
        locale="en",
        components=[component_block, tda_matched_block, unmatched_block],
        execution_id="exe_custom_export_123",
    )

    assert payload.filename == "execution_export_exe_custom_export_123.xlsx"
    assert len(payload.content_bytes) > 0


@pytest.mark.asyncio
async def test_export_flat_csv_default_execution_id() -> None:
    """Test flat CSV export without explicit execution_id."""
    service = ExportService()
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED)
    report_dto = _build_sample_report_dto()
    matching_matrix = MatrixScorecardRowDTO(
        block_id="blk_0123456789abcdef",
        name="Axis 1",
        label_i18n=I18nText(translations={"fi": "Akseli 1", "en": "Axis 1"}),
        score=4.5,
        scale_max=5.0,
        row_explanation="Explanation for row 1",
        is_evaluative=True,
    )
    matrices = [matching_matrix]

    payload = await service.export_flat_csv(
        execution=exec_record,
        report_dto=report_dto,
        matrices=matrices,
        locale="fi",
    )

    assert payload.filename == "execution_export_exe_0123456789abcdef.csv"
    assert len(payload.content_bytes) > 0
    csv_text = payload.content_bytes.decode("utf-8-sig")
    assert "Logiikkamatriisi" in csv_text


@pytest.mark.asyncio
async def test_export_excel_fails_when_matrices_empty() -> None:
    """Test that empty matrices list triggers Fail-Fast AppException(400, VALIDATION_FAILED)."""
    service = ExportService()
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED)
    report_dto = _build_sample_report_dto()

    with pytest.raises(AppException) as exc_info:
        await service.export_excel(
            execution=exec_record,
            report_dto=report_dto,
            matrices=[],
        )

    assert exc_info.value.status_code == 400
    assert exc_info.value.details["error_code"] == ErrorCodes.VALIDATION_FAILED.value
    assert "Execution has no evaluative matrices for Excel export." in exc_info.value.message


@pytest.mark.asyncio
async def test_export_excel_emits_exact_9_rows_summary_and_claim_classifications() -> None:
    """Test that export_excel emits exactly 9 matrix summary rows and explicit claim classifications."""
    import io

    import pandas as pd

    service = ExportService()
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED)

    matrices = _build_sample_matrix_rows(9)

    atom_positive = AtomResultDTO(
        tda_id="tda_pos_001",
        matrix_id="blk_matrix_0001",
        status=ExecutionStatus.PASSED,
        source_quotes="Positive quote",
        evaluation_reasoning="Demonstrated positive competence.",
        is_inverse_evidence=False,
    )
    atom_inverse = AtomResultDTO(
        tda_id="tda_inv_001",
        matrix_id="blk_matrix_0002",
        status=ExecutionStatus.PASSED,
        source_quotes=None,
        evaluation_reasoning="No fallacy detected.",
        is_inverse_evidence=True,
    )

    report_dto = ReportDataDTO(
        workflow_id="wor_0123456789abcdef",
        execution_id="exe_0123456789abcdef",
        profile_id="prof_0123456789abcdef",
        global_score=4.0,
        has_warning=False,
        inner_sdui_blocks=[],
        results=[atom_positive, atom_inverse],
        hydrated_references={
            "tda_pos_001": HydratedAtomDTO(
                sdui_component=LaxSDUIComponentType.BOOLEAN_CARD,
                resolved_claim="Competence claim",
            ),
            "tda_inv_001": HydratedAtomDTO(
                sdui_component=LaxSDUIComponentType.BOOLEAN_CARD,
                resolved_claim="Fallacy absence claim",
            ),
        },
    )

    payload = await service.export_excel(
        execution=exec_record,
        report_dto=report_dto,
        matrices=matrices,
        locale="fi",
    )

    excel_file = io.BytesIO(payload.content_bytes)
    summary_df = pd.read_excel(excel_file, sheet_name="Yhteenveto")
    raw_df = pd.read_excel(excel_file, sheet_name="Raakadata")

    assert len(summary_df) == 9
    assert len(summary_df) == len(matrices)
    for i in range(len(summary_df)):
        assert summary_df["Logiikkamatriisi"].iloc[i] == f"Matriisi {i + 1}"
        assert summary_df["Pisteet (raaka)"].iloc[i] == 4.0
        assert summary_df["Asteikon maksimi"].iloc[i] == 5.0

    assert len(raw_df) == 2
    assert len(raw_df) == len(report_dto.results)
    assert "Väitetyyppi" in raw_df.columns
    assert "Tulos (Status)" in raw_df.columns
    assert "Arviointikriteeri" in raw_df.columns

    pos_row = raw_df[raw_df["Arviointikriteeri"] == "Competence claim"].iloc[0]
    assert pos_row["Väitetyyppi"] == "Positiivinen kyvykkyys"
    assert pos_row["Tulos (Status)"] == 1

    inv_row = raw_df[raw_df["Arviointikriteeri"] == "Fallacy absence claim"].iloc[0]
    assert inv_row["Väitetyyppi"] == "Virhedetektori / Anti-pattern"
    assert inv_row["Tulos (Status)"] == 1
    assert "Falsifiointi" not in raw_df.columns
    assert "Sisäistetty sääntö" not in raw_df.columns


@pytest.mark.asyncio
async def test_export_excel_with_prompt_block_repo_resolves_matrix_names_and_inverse_claims_positive() -> None:
    prompt_block_repo = InMemoryPromptBlockRepository()
    lbl = I18nText(
        translations={
            "fi": "Aktiivinen ohjaus (Performatiivisuus ja Goodhartin Laki)",
            "en": "Active Steering (Performativity and Goodhart's Law)",
        }
    )
    desc = I18nText(translations={"fi": "Kuvaus", "en": "Description"})
    scale = MatrixScale(score=1, ai_label="L1", claims=[])
    matrix_block = MatrixPromptBlock(
        id="blk_53f32679aa514fcb",
        slug="active_steering",
        label=lbl,
        description=desc,
        scales=[scale],
        ai_description="Operational rule for active steering",
    )
    await prompt_block_repo.create_prompt_block(matrix_block)

    service = ExportService(prompt_block_repo=prompt_block_repo)
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED)
    matrices = _build_sample_matrix_rows(1)

    atom_inv = AtomResultDTO(
        tda_id="tda_53f32679aa514fcb0000000000000000",
        matrix_id="blk_53f32679aa514fcb",
        status=ExecutionStatus.PASSED,
        source_quotes=None,
        evaluation_reasoning="No performativity error detected in leadership steering.",
        is_inverse_evidence=True,
    )

    report_dto = ReportDataDTO(
        workflow_id="wor_0123456789abcdef",
        execution_id="exe_0123456789abcdef",
        profile_id="prof_0123456789abcdef",
        global_score=4.0,
        has_warning=False,
        inner_sdui_blocks=[],
        results=[atom_inv],
        hydrated_references={
            "tda_53f32679aa514fcb0000000000000000": HydratedAtomDTO(
                sdui_component=LaxSDUIComponentType.BOOLEAN_CARD,
                resolved_claim="Goodhartin laki vältetty",
            )
        },
    )

    payload = await service.export_excel(
        execution=exec_record,
        report_dto=report_dto,
        matrices=matrices,
        locale="fi",
    )

    raw_df = pd.read_excel(io.BytesIO(payload.content_bytes), sheet_name="Raakadata")
    assert len(raw_df) == 1
    row = raw_df.iloc[0]
    atom_col_name = ReportHeaderResolver.get_atom_column_header(ReportAtomColumn.MATRIX, "fi")
    assert row[atom_col_name] == "Aktiivinen ohjaus (Performatiivisuus ja Goodhartin Laki)"
    assert row["Väitetyyppi"] == "Virhedetektori / Anti-pattern"
    assert row["Tulos (Status)"] == 1
    assert prompt_block_repo._call_counts["get_all_prompt_blocks"] == 1


@pytest.mark.asyncio
async def test_export_excel_with_failed_inverse_claim_emits_status_zero_negative() -> None:
    prompt_block_repo = InMemoryPromptBlockRepository()
    tda_id = "tda_99999999aaaaaaaa0000000000000000"
    tda = TDAAssertion(
        tda_id=tda_id,
        inverse_evidence=True,
        aggregation_mode="EXISTS",
        concept_description="Detect steering error or performative bias",
    )
    claim = MatrixClaim(
        label=I18nText(translations={"fi": "Virhevapaa ohjaus", "en": "Error-free steering"}),
        tda_assertions=[tda],
    )
    scale = MatrixScale(score=1, ai_label="L1", claims=[claim])
    matrix_block = MatrixPromptBlock(
        id="blk_0000000000000001",
        slug="matrix_0001",
        label=I18nText(translations={"fi": "Matriisi 1", "en": "Matrix 1"}),
        description=I18nText(translations={"fi": "Kuvaus 1", "en": "Desc 1"}),
        scales=[scale],
        ai_description="Operational rule for matrix 1",
    )
    await prompt_block_repo.create_prompt_block(matrix_block)

    service = ExportService(prompt_block_repo=prompt_block_repo)
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED)
    matrices = [
        MatrixScorecardRowDTO(
            block_id="blk_0000000000000001",
            name="Matrix 1",
            label_i18n=I18nText(translations={"fi": "Matriisi 1", "en": "Matrix 1"}),
            score=4.0,
            scale_max=5.0,
            row_explanation="Row explanation 1",
            is_evaluative=True,
        )
    ]

    atom_failed = AtomResultDTO(
        tda_id=tda_id,
        matrix_id="blk_0000000000000001",
        status=ExecutionStatus.FAILED,
        source_quotes=None,
        evaluation_reasoning="Performative bias was detected in steering.",
        is_inverse_evidence=False,
    )

    report_dto = ReportDataDTO(
        workflow_id="wor_0123456789abcdef",
        execution_id="exe_0123456789abcdef",
        profile_id="prof_0123456789abcdef",
        global_score=4.0,
        has_warning=False,
        inner_sdui_blocks=[],
        results=[atom_failed],
        hydrated_references={
            tda_id: HydratedAtomDTO(
                sdui_component=LaxSDUIComponentType.BOOLEAN_CARD,
                resolved_claim="Virhevapaa ohjaus",
            )
        },
    )

    payload = await service.export_excel(
        execution=exec_record,
        report_dto=report_dto,
        matrices=matrices,
        locale="fi",
    )

    raw_df = pd.read_excel(io.BytesIO(payload.content_bytes), sheet_name="Raakadata")
    assert len(raw_df) == 1
    row = raw_df.iloc[0]
    assert row["Väitetyyppi"] == "Virhedetektori / Anti-pattern"
    assert row["Tulos (Status)"] == 0


@pytest.mark.asyncio
async def test_export_excel_missing_matrix_in_repo_raises_fail_fast_negative() -> None:
    prompt_block_repo = InMemoryPromptBlockRepository()
    service = ExportService(prompt_block_repo=prompt_block_repo)
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED)
    matrices = _build_sample_matrix_rows(1)

    unknown_atom = AtomResultDTO(
        tda_id="tda_0123456789abcdef0000000000000000",
        matrix_id="blk_unknown_99999999",
        status=ExecutionStatus.PASSED,
        source_quotes="Some valid quote",
        evaluation_reasoning="Evaluation succeeded.",
    )

    report_dto = ReportDataDTO(
        workflow_id="wor_0123456789abcdef",
        execution_id="exe_0123456789abcdef",
        profile_id="prof_0123456789abcdef",
        global_score=4.0,
        has_warning=False,
        inner_sdui_blocks=[],
        results=[unknown_atom],
        hydrated_references={
            "tda_0123456789abcdef0000000000000000": HydratedAtomDTO(
                sdui_component=LaxSDUIComponentType.BOOLEAN_CARD,
                resolved_claim="Some claim",
            )
        },
    )

    with pytest.raises(AppException) as exc_info:
        await service.export_excel(
            execution=exec_record,
            report_dto=report_dto,
            matrices=matrices,
            locale="fi",
        )

    assert exc_info.value.status_code == 404
    assert exc_info.value.details["error_code"] == ErrorCodes.RESOURCE_NOT_FOUND.value
    assert "Strict Fail-Fast: Matrix block 'blk_unknown_99999999'" in exc_info.value.message


@pytest.mark.asyncio
async def test_export_excel_tab1_ssot_localization() -> None:
    """Verify Tab 1 columns match ReportHeaderResolver SSOT headers in both fi and en."""
    service = ExportService()
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED)
    report_dto = _build_sample_report_dto()
    matrices = [
        MatrixScorecardRowDTO(
            block_id="blk_0123456789abcdef",
            name="Axis 1",
            label_i18n=I18nText(translations={"fi": "Akseli 1", "en": "Axis 1"}),
            score=4.5,
            scale_max=5.0,
            row_explanation="Explanation for row 1",
            is_evaluative=True,
        ),
        MatrixScorecardRowDTO(
            block_id="blk_matrix_0002",
            name="Matrix 2",
            label_i18n=I18nText(translations={"fi": "Matriisi 2", "en": "Matrix 2"}),
            score=4.0,
            scale_max=5.0,
            row_explanation="Explanation for row 2",
            is_evaluative=True,
        ),
    ]

    for locale in ["fi", "en"]:
        payload = await service.export_excel(
            execution=exec_record,
            report_dto=report_dto,
            matrices=matrices,
            locale=locale,
        )
        sheet_name = ReportHeaderResolver.get_sheet_name(ReportSheetKey.SUMMARY, locale=locale)
        df = pd.read_excel(io.BytesIO(payload.content_bytes), sheet_name=sheet_name)
        expected_headers = list(ReportHeaderResolver.get_all_matrix_headers(locale=locale).values())
        assert list(df.columns) == expected_headers
        assert len(df) == len(matrices)


@pytest.mark.asyncio
async def test_export_excel_tab2_ssot_localization() -> None:
    """Verify Tab 2 headers match ReportAtomColumn SSOT headers and result_status is strictly binary."""
    service = ExportService()
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED)
    report_dto = _build_sample_report_dto()
    matching_matrix = MatrixScorecardRowDTO(
        block_id="blk_0123456789abcdef",
        name="Axis 1",
        label_i18n=I18nText(translations={"fi": "Akseli 1", "en": "Axis 1"}),
        score=4.5,
        scale_max=5.0,
        row_explanation="Explanation for row 1",
        is_evaluative=True,
    )
    matrices = [matching_matrix]

    for locale in ["fi", "en"]:
        payload = await service.export_excel(
            execution=exec_record,
            report_dto=report_dto,
            matrices=matrices,
            locale=locale,
        )
        sheet_name = ReportHeaderResolver.get_sheet_name(ReportSheetKey.RAW_DATA, locale=locale)
        df = pd.read_excel(io.BytesIO(payload.content_bytes), sheet_name=sheet_name)
        expected_headers = ReportHeaderResolver.get_all_atom_headers(locale=locale)
        assert list(df.columns) == expected_headers
        assert len(df) == len(report_dto.results)
        status_header = ReportHeaderResolver.get_atom_column_header(ReportAtomColumn.RESULT_STATUS, locale=locale)
        for status_val in df[status_header]:
            assert status_val in (0, 1)


@pytest.mark.asyncio
async def test_export_flat_csv_fails_non_passed_status() -> None:
    service = ExportService()
    exec_record = _build_sample_execution(status=ExecutionStatus.FAILED)
    with pytest.raises(AppException) as exc_info:
        await service.export_flat_csv(execution=exec_record, report_dto=None, matrices=[])
    assert exc_info.value.status_code == 400
    assert "Execution must be in PASSED state" in exc_info.value.message


@pytest.mark.asyncio
async def test_export_flat_csv_fails_when_no_atoms() -> None:
    service = ExportService()
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED, has_atoms=False)
    report_dto = _build_sample_report_dto(has_atoms=False)
    with pytest.raises(AppException) as exc_info:
        await service.export_flat_csv(execution=exec_record, report_dto=report_dto, matrices=[])
    assert exc_info.value.status_code == 400
    assert "Execution has no scoreable atoms" in exc_info.value.message


@pytest.mark.asyncio
async def test_export_excel_with_scorecard_atoms_and_rich_matrix_metadata() -> None:
    service = ExportService()
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED)

    scorecard_atom = ScorecardAtomDTO(
        atom_id="tda_eval_001",
        claim_label="Claim label from scorecard",
        level=2,
        level_name="Level Two",
        extracted_facts={},
        exact_quotes=[],
        internal_logic_en=ReasoningStepDTO(
            step_1_identify_premise="Premise",
            step_2_scan_source="Scan",
            step_3_evaluate_anti_patterns="Anti-pattern",
            step_4_final_conclusion="Conclusion",
        ),
        status=ExecutionStatus.PASSED,
        semantic_reasoning="Semantic reason",
        contextual_override=False,
        chart_display_label="Chart display label",
        visual_intent=VisualIntent.NEUTRAL,
    )
    rich_matrix = MatrixScorecardRowDTO(
        block_id="blk_rich_0001",
        name="Rich Matrix",
        label_i18n=I18nText(translations={"fi": "Rikas matriisi", "en": "Rich Matrix"}),
        context_target="ctx_tgt_plain",
        context_target_label=I18nText(translations={"fi": "Kohteen nimike", "en": "Target title"}),
        score=3.5,
        scale_max=None,
        normalized_score=70.0,
        level_breakdown={"L1": "1", "L2": "2"},
        true_atoms=3,
        total_atoms=4,
        cited_text_quote="Direct citation quote from matrix",
        cited_source_title="Source Title Doc",
        cited_source_id="src_doc_01",
        row_explanation="Explanation for rich matrix",
        evaluated_atoms=[scorecard_atom],
        is_evaluative=True,
    )

    atom_without_matrix_id = AtomResultDTO(
        tda_id="tda_eval_001",
        matrix_id=None,
        status=ExecutionStatus.PASSED,
        source_quotes="Verbatim quote for atom without matrix_id",
        evaluation_reasoning="Reasoning for atom without matrix_id",
    )
    atom_with_ref_quote = AtomResultDTO(
        tda_id="tda_eval_002",
        matrix_id="blk_rich_0001",
        status=ExecutionStatus.PASSED,
        source_quotes=None,
        evaluation_reasoning="Reasoning with ref quote fallback",
        is_inverse_evidence=True,
    )

    report_dto = ReportDataDTO(
        workflow_id="wor_0123456789abcdef",
        execution_id="exe_0123456789abcdef",
        profile_id="prof_0123456789abcdef",
        global_score=3.5,
        has_warning=False,
        inner_sdui_blocks=[],
        results=[atom_without_matrix_id, atom_with_ref_quote],
        hydrated_references={
            "tda_eval_001": HydratedAtomDTO(
                sdui_component=LaxSDUIComponentType.BOOLEAN_CARD,
                resolved_claim="Claim for eval 1",
            ),
            "tda_eval_002": HydratedAtomDTO(
                sdui_component=LaxSDUIComponentType.BOOLEAN_CARD,
                resolved_claim="Claim for eval 2",
                source_quotes="Hydrated fallback quote",
            ),
        },
    )

    payload = await service.export_excel(
        execution=exec_record,
        report_dto=report_dto,
        matrices=[rich_matrix],
        locale="fi",
    )
    assert len(payload.content_bytes) > 0


@pytest.mark.asyncio
async def test_export_excel_atom_missing_matrix_id_and_parent_fails_fast() -> None:
    service = ExportService()
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED)
    orphan_atom = AtomResultDTO(
        tda_id="tda_orphan_999",
        matrix_id=None,
        status=ExecutionStatus.PASSED,
        source_quotes="Quote",
        evaluation_reasoning="Reason",
    )
    report_dto = ReportDataDTO(
        workflow_id="wor_0123456789abcdef",
        execution_id="exe_0123456789abcdef",
        profile_id="prof_0123456789abcdef",
        global_score=3.5,
        has_warning=False,
        inner_sdui_blocks=[],
        results=[orphan_atom],
        hydrated_references={
            "tda_orphan_999": HydratedAtomDTO(
                sdui_component=LaxSDUIComponentType.BOOLEAN_CARD,
                resolved_claim="Orphan claim",
            )
        },
    )
    rich_matrix = MatrixScorecardRowDTO(
        block_id="blk_rich_0001",
        name="Rich Matrix",
        label_i18n=I18nText(translations={"fi": "Rikas matriisi", "en": "Rich Matrix"}),
        score=3.5,
        scale_max=5.0,
        row_explanation="Explanation",
        evaluated_atoms=[],
        is_evaluative=True,
    )
    with pytest.raises(AppException) as exc_info:
        await service.export_excel(
            execution=exec_record,
            report_dto=report_dto,
            matrices=[rich_matrix],
            locale="fi",
        )
    assert exc_info.value.status_code == 404
    assert "Strict Fail-Fast: Matrix block for atom 'tda_orphan_999' not found" in exc_info.value.message


@pytest.mark.asyncio
async def test_export_flat_csv_resolves_matrix_from_prompt_block_repo() -> None:
    """Regression test: export_flat_csv must resolve matrix blocks from prompt_block_repo.

    When an atom belongs to an informational or secondary matrix block not included
    in the evaluative 'matrices' list (e.g. blk_53f32679aa514fcb), export_flat_csv
    must resolve the matrix metadata instead of crashing with 404 RESOURCE_NOT_FOUND.
    """
    prompt_block_repo = InMemoryPromptBlockRepository()
    lbl = I18nText(translations={"fi": "Aktiivinen ohjaus", "en": "Active Control"})
    tda_assertion = TDAAssertion(
        tda_id="tda_135ee4d2f4e28b3305cbae57482e34b0",
        inverse_evidence=False,
        aggregation_mode="ALL_MUST_COMPLY",
        concept_description="Active steering assertion",
    )
    claim = MatrixClaim(
        label=I18nText(translations={"fi": "Väite 1", "en": "Claim 1"}),
        tda_assertions=[tda_assertion],
    )
    scale = MatrixScale(score=1, ai_label="L1", claims=[claim])
    desc = I18nText(translations={"fi": "Kuvaus", "en": "Description"})
    informational_block = MatrixPromptBlock(
        id="blk_53f32679aa514fcb",
        slug="m_slug",
        label=lbl,
        description=desc,
        scales=[scale],
        ai_description="English operational rule",
        is_evaluative=False,
    )
    await prompt_block_repo.create_prompt_block(informational_block)
    service = ExportService(prompt_block_repo=prompt_block_repo)
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED)
    atom_result = AtomResultDTO(
        tda_id="tda_135ee4d2f4e28b3305cbae57482e34b0",
        matrix_id="blk_53f32679aa514fcb",
        status=ExecutionStatus.PASSED,
        source_quotes="Verbatim quote",
        evaluation_reasoning="Evaluation reasoning",
    )
    report_dto = ReportDataDTO(
        workflow_id="wor_0123456789abcdef",
        execution_id="exe_0123456789abcdef",
        profile_id="prof_0123456789abcdef",
        global_score=4.0,
        has_warning=False,
        inner_sdui_blocks=[],
        results=[atom_result],
        hydrated_references={
            "tda_135ee4d2f4e28b3305cbae57482e34b0": HydratedAtomDTO(
                sdui_component=LaxSDUIComponentType.BOOLEAN_CARD,
                resolved_claim="Väite 1",
            )
        },
    )
    evaluative_matrix = MatrixScorecardRowDTO(
        block_id="blk_eval_0001",
        name="Evaluative Matrix",
        label_i18n=I18nText(translations={"fi": "Arvioiva", "en": "Evaluative"}),
        score=4.0,
        scale_max=5.0,
        row_explanation="Explanation",
        is_evaluative=True,
    )
    payload = await service.export_flat_csv(
        execution=exec_record,
        report_dto=report_dto,
        matrices=[evaluative_matrix],
        locale="fi",
    )
    assert payload is not None
    csv_text = payload.content_bytes.decode("utf-8-sig")
    assert "Aktiivinen ohjaus" in csv_text
    assert "L1" in csv_text


@pytest.mark.asyncio
async def test_informational_matrix_parity_excel_and_csv() -> None:
    """Verify informational matrices (is_evaluative=False) have complete parity in Excel and CSV.

    Scores, levels, level names, and atom rows are computed and exported identically,
    while global_score average remains driven solely by evaluative matrices.
    """
    prompt_block_repo = InMemoryPromptBlockRepository()
    info_block = MatrixPromptBlock(
        id="blk_1111222233334444",
        slug="info_slug",
        label=I18nText(translations={"fi": "Informaatiomatriisi", "en": "Informational Matrix"}),
        description=I18nText(translations={"fi": "Kuvaus", "en": "Desc"}),
        scales=[
            MatrixScale(
                score=3,
                ai_label="Taso 3",
                name=I18nText(translations={"fi": "Taso 3 Nimi", "en": "Level 3 Name"}),
                claims=[],
            )
        ],
        ai_description="Info rule",
        is_evaluative=False,
    )
    await prompt_block_repo.create_prompt_block(info_block)

    service = ExportService(prompt_block_repo=prompt_block_repo)
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED)

    info_scorecard_atom = ScorecardAtomDTO(
        atom_id="tda_1111222233334444",
        claim_label="Informaatioväite",
        level=3,
        level_name="Taso 3 Nimi",
        extracted_facts={},
        exact_quotes=[],
        internal_logic_en=ReasoningStepDTO(
            step_1_identify_premise="Premise",
            step_2_scan_source="Scan",
            step_3_evaluate_anti_patterns="Anti-pattern",
            step_4_final_conclusion="Conclusion",
        ),
        status=ExecutionStatus.PASSED,
        semantic_reasoning="Reasoning for info atom",
        contextual_override=False,
        chart_display_label="Chart label",
        visual_intent=VisualIntent.NEUTRAL,
    )

    eval_scorecard_atom = ScorecardAtomDTO(
        atom_id="tda_5555666677778888",
        claim_label="Arviointiväite",
        level=5,
        level_name="Taso 5 Nimi",
        extracted_facts={},
        exact_quotes=[],
        internal_logic_en=ReasoningStepDTO(
            step_1_identify_premise="Premise",
            step_2_scan_source="Scan",
            step_3_evaluate_anti_patterns="Anti-pattern",
            step_4_final_conclusion="Conclusion",
        ),
        status=ExecutionStatus.PASSED,
        semantic_reasoning="Reasoning for eval atom",
        contextual_override=False,
        chart_display_label="Chart label",
        visual_intent=VisualIntent.NEUTRAL,
    )

    eval_matrix = MatrixScorecardRowDTO(
        block_id="blk_5555666677778888",
        name="Eval Matrix",
        label_i18n=I18nText(translations={"fi": "Arvioiva Matriisi", "en": "Evaluative Matrix"}),
        score=5.0,
        scale_max=5.0,
        normalized_score=100.0,
        level_breakdown={"L5": "1"},
        row_explanation="Evaluative explanation",
        evaluated_atoms=[eval_scorecard_atom],
        is_evaluative=True,
    )

    info_matrix = MatrixScorecardRowDTO(
        block_id="blk_1111222233334444",
        name="Info Matrix",
        label_i18n=I18nText(translations={"fi": "Informaatiomatriisi", "en": "Informational Matrix"}),
        score=3.0,
        scale_max=5.0,
        normalized_score=60.0,
        level_breakdown={"L3": "1"},
        row_explanation="Informational explanation",
        evaluated_atoms=[info_scorecard_atom],
        is_evaluative=False,
    )

    atom_eval = AtomResultDTO(
        tda_id="tda_5555666677778888",
        matrix_id="blk_5555666677778888",
        status=ExecutionStatus.PASSED,
        source_quotes="Eval quote",
        evaluation_reasoning="Eval reasoning",
    )
    atom_info = AtomResultDTO(
        tda_id="tda_1111222233334444",
        matrix_id="blk_1111222233334444",
        status=ExecutionStatus.PASSED,
        source_quotes="Info quote",
        evaluation_reasoning="Info reasoning",
    )

    report_dto = ReportDataDTO(
        workflow_id="wor_0123456789abcdef",
        execution_id="exe_0123456789abcdef",
        profile_id="prof_0123456789abcdef",
        global_score=5.0,
        has_warning=False,
        inner_sdui_blocks=[],
        results=[atom_eval, atom_info],
        hydrated_references={
            "tda_5555666677778888": HydratedAtomDTO(
                sdui_component=LaxSDUIComponentType.BOOLEAN_CARD,
                resolved_claim="Arviointiväite",
            ),
            "tda_1111222233334444": HydratedAtomDTO(
                sdui_component=LaxSDUIComponentType.BOOLEAN_CARD,
                resolved_claim="Informaatioväite",
            ),
        },
    )

    all_matrices = [eval_matrix, info_matrix]

    excel_payload = await service.export_excel(
        execution=exec_record,
        report_dto=report_dto,
        matrices=all_matrices,
        locale="fi",
        all_matrices=all_matrices,
    )
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(excel_payload.content_bytes))
    sheet1 = wb[LocalizationService.translate("export_sheet_summary", "fi")]
    sheet1_rows = list(sheet1.iter_rows(values_only=True))
    assert len(sheet1_rows) == 3
    matrix_names = [row[1] for row in sheet1_rows[1:]]
    assert "Arvioiva Matriisi" in matrix_names
    assert "Informaatiomatriisi" in matrix_names

    sheet2 = wb[LocalizationService.translate("export_sheet_raw_data", "fi")]
    sheet2_rows = list(sheet2.iter_rows(values_only=True))
    assert len(sheet2_rows) == 3
    raw_levels = [row[2] for row in sheet2_rows[1:]]
    assert 5 in raw_levels
    assert 3 in raw_levels
    raw_level_names = [row[3] for row in sheet2_rows[1:]]
    assert "Taso 5 Nimi" in raw_level_names
    assert "Taso 3 Nimi" in raw_level_names

    csv_payload = await service.export_flat_csv(
        execution=exec_record,
        report_dto=report_dto,
        matrices=all_matrices,
        locale="fi",
        all_matrices=all_matrices,
    )
    csv_text = csv_payload.content_bytes.decode("utf-8-sig")
    assert "Arvioiva Matriisi" in csv_text
    assert "Informaatiomatriisi" in csv_text
    assert "Taso 5 Nimi" in csv_text
    assert "Taso 3 Nimi" in csv_text


def test_report_atom_column_has_14_members() -> None:
    """Verify ReportAtomColumn has exactly 14 members including causal fields and ML columns."""
    assert len(ReportAtomColumn) == 14
    assert ReportAtomColumn.IS_INVERSE.value == "is_inverse"
    assert ReportAtomColumn.POINTS_AWARDED.value == "points_awarded"
    assert ReportAtomColumn.CAUSAL_STATUS.value == "causal_status"
    assert ReportAtomColumn.BLAME_PARENT_ID.value == "blame_parent_id"
    assert ReportAtomColumn.DEPENDENT_COUNT.value == "dependent_count"


def test_denormalized_csv_headers_has_22_columns() -> None:
    """Verify ReportHeaderResolver.get_all_denormalized_csv_headers has 22 columns."""
    headers_en = ReportHeaderResolver.get_all_denormalized_csv_headers("en")
    assert len(headers_en) == 22
    assert headers_en[-3:] == ["Causal Status", "Root Cause ID", "Dependent Count"]

    headers_fi = ReportHeaderResolver.get_all_denormalized_csv_headers("fi")
    assert len(headers_fi) == 22
    assert headers_fi[-3:] == ["Kausaalistatus", "Juurisyyn tunnus", "Riippuvuuksien lkm"]


@pytest.mark.asyncio
async def test_export_flat_csv_row_count_matches_results() -> None:
    """Verify flat CSV export streams exactly 1 header + N atom rows, zero synthetic summary rows."""
    service = ExportService()
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED)
    report_dto = _build_sample_report_dto()
    matrices = [
        MatrixScorecardRowDTO(
            block_id="blk_0123456789abcdef",
            name="Matrix 1",
            label_i18n=I18nText(translations={"en": "Matrix 1"}),
            score=4.0,
            scale_max=5.0,
            row_explanation="Row exp",
            is_evaluative=True,
        )
    ]
    payload = await service.export_flat_csv(
        execution=exec_record,
        report_dto=report_dto,
        matrices=matrices,
        locale="en",
    )
    csv_text = payload.content_bytes.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(csv_text)))
    assert len(rows) == len(report_dto.results) + 1
    assert len(rows[0]) == 22


@pytest.mark.asyncio
async def test_export_causal_enrichment_parity() -> None:
    """Verify Excel Tab 2 and CSV stream rows contain populated causal fields matching graph."""
    service = ExportService()

    node1 = CausalNodeDTO(
        node_id="tda_0000000000000001",
        claim="Foundational architecture assumption holds true.",
        source_quotes="Foundational verified quote.",
        paragraph_ref="Section 1",
        status=ExecutionStatus.PASSED,
        blame_parent_ids=[],
        dependent_child_ids=["tda_0000000000000002"],
        tda_score=5.0,
        tda_level_key="L5",
    )
    node2 = CausalNodeDTO(
        node_id="tda_0000000000000002",
        claim="Downstream interface verification.",
        source_quotes="Interface failure quote.",
        paragraph_ref="Section 2",
        status=ExecutionStatus.FAILED,
        blame_parent_ids=["tda_0000000000000001"],
        dependent_child_ids=[],
        tda_score=1.0,
        tda_level_key="L1",
    )
    edge = CausalEdgeDTO(
        source_id="tda_0000000000000001",
        target_id="tda_0000000000000002",
        reasoning="Dependency implies interface requirement.",
    )
    graph = CausalGraphPayloadDTO(
        nodes=[node1, node2],
        edges=[edge],
        root_cause_node_ids=["tda_0000000000000001"],
        cycle_detected=False,
        isolated_node_ids=[],
    )
    diag = CausalRootCauseDiagnosisDTO(
        failing_node_id="tda_0000000000000002",
        root_cause_parent_id="tda_0000000000000001",
        paragraph_ref="Section 2",
        explanation="Root assumption lacked schema boundary.",
        penalty_points=1.5,
    )
    audit = AntiFluffAuditDTO(
        isolated_concept_count=0,
        cycle_detected=False,
        unsupported_claims_count=1,
        logical_cohesion_score=0.9,
        detected_hollow_terms=[],
    )
    remediation = PrescriptiveRemediationDTO(
        target_node_id="tda_0000000000000002",
        required_action="Add extra='forbid' to model.",
        unlocks_node_ids=["tda_0000000000000003"],
        potential_score_impact=1.5,
    )
    fair_scoring = FairScoringBreakdownDTO(
        independent_error_count=1,
        cascading_fault_count=1,
        raw_penalty_points=3.0,
        deduplicated_penalty_points=1.5,
        dampened_savings_points=1.5,
    )
    fusion_result = CausalTdaFusionResultDTO(
        graph=graph,
        root_cause_diagnoses=[diag],
        anti_fluff_audit=audit,
        prescriptive_remediations=[remediation],
        fair_scoring=fair_scoring,
    )

    exec_record = ExecutionRecord(
        id="exe_0123456789abcdef",
        workflow_id="wor_0123456789abcdef",
        organization_id="org_0123456789abcdef",
        target_locale="en",
        status=ExecutionStatus.PASSED,
        execution_trace=[
            TraceEvent(
                step_name="causal_discovery",
                event_type="output",
                content=fusion_result,
            )
        ],
    )

    atom1 = AtomResultDTO(
        tda_id="tda_0000000000000001",
        matrix_id="blk_0123456789abcdef",
        status=ExecutionStatus.PASSED,
        source_quotes="Foundational verified quote.",
        evaluation_reasoning="Valid reasoning 1",
    )
    atom2 = AtomResultDTO(
        tda_id="tda_0000000000000002",
        matrix_id="blk_0123456789abcdef",
        status=ExecutionStatus.FAILED,
        source_quotes=None,
        evaluation_reasoning="Failure reasoning 2",
    )
    report_dto = ReportDataDTO(
        workflow_id="wor_0123456789abcdef",
        execution_id="exe_0123456789abcdef",
        profile_id="prf_0123456789abcdef",
        results=[atom1, atom2],
        hydrated_references={
            "tda_0000000000000001": HydratedAtomDTO(
                sdui_component=LaxSDUIComponentType.BOOLEAN_CARD,
                resolved_claim="Claim 1",
            ),
            "tda_0000000000000002": HydratedAtomDTO(
                sdui_component=LaxSDUIComponentType.BOOLEAN_CARD,
                resolved_claim="Claim 2",
            ),
        },
    )
    matrices = [
        MatrixScorecardRowDTO(
            block_id="blk_0123456789abcdef",
            name="Matrix 1",
            label_i18n=I18nText(translations={"en": "Matrix 1"}),
            score=4.0,
            scale_max=5.0,
            row_explanation="Explanation",
            is_evaluative=True,
        )
    ]

    # 1. Verify Excel Tab 2
    excel_payload = await service.export_excel(
        execution=exec_record,
        report_dto=report_dto,
        matrices=matrices,
        locale="en",
    )
    sheet_name = ReportHeaderResolver.get_sheet_name(ReportSheetKey.RAW_DATA, locale="en")
    df = pd.read_excel(io.BytesIO(excel_payload.content_bytes), sheet_name=sheet_name)
    assert len(df) == 2
    assert list(df.columns) == ReportHeaderResolver.get_all_atom_headers("en")

    causal_status_col = ReportHeaderResolver.get_atom_column_header(ReportAtomColumn.CAUSAL_STATUS, "en")
    blame_parent_col = ReportHeaderResolver.get_atom_column_header(ReportAtomColumn.BLAME_PARENT_ID, "en")
    dep_count_col = ReportHeaderResolver.get_atom_column_header(ReportAtomColumn.DEPENDENT_COUNT, "en")

    # Node 1: validated, no blame parent, 1 dependent child
    assert df.iloc[0][causal_status_col] == "validated"
    assert pd.isna(df.iloc[0][blame_parent_col]) or df.iloc[0][blame_parent_col] == ""
    assert int(df.iloc[0][dep_count_col]) == 1

    # Node 2: cascading_fault, blame parent 'tda_0000000000000001', 0 dependent children
    assert df.iloc[1][causal_status_col] == "cascading_fault"
    assert df.iloc[1][blame_parent_col] == "tda_0000000000000001"
    assert int(df.iloc[1][dep_count_col]) == 0

    # 2. Verify CSV stream
    csv_payload = await service.export_flat_csv(
        execution=exec_record,
        report_dto=report_dto,
        matrices=matrices,
        locale="en",
    )
    csv_rows = list(csv.reader(io.StringIO(csv_payload.content_bytes.decode("utf-8-sig"))))
    assert len(csv_rows) == 3  # 1 header + 2 atom rows
    assert csv_rows[0] == ReportHeaderResolver.get_all_denormalized_csv_headers("en")

    # Node 1 row
    assert csv_rows[1][19] == "validated"
    assert csv_rows[1][20] == ""
    assert csv_rows[1][21] == "1"

    # Node 2 row
    assert csv_rows[2][19] == "cascading_fault"
    assert csv_rows[2][20] == "tda_0000000000000001"
    assert csv_rows[2][21] == "0"


def test_export_forensic_atom_dto_extra_field_raises_validation_error() -> None:
    """Error path: ExportForensicAtomDTO raises ValidationError when unknown field is passed."""
    with pytest.raises(ValidationError):
        ExportForensicAtomDTO.model_validate(
            {
                "matrix_label": "M",
                "context_target": "C",
                "level": 1,
                "level_name": "L1",
                "criterion": "Crit",
                "claim_type": "Claim",
                "result_status": 1,
                "quotes": "Quotes",
                "ai_reasoning": "Reasoning",
                "unknown_extra_field": "disallowed",
            }
        )


def test_export_denormalized_flat_row_dto_extra_field_raises_validation_error() -> None:
    """Error path: ExportDenormalizedFlatRowDTO raises ValidationError when unknown field is passed."""
    with pytest.raises(ValidationError):
        ExportDenormalizedFlatRowDTO.model_validate(
            {
                "execution_id": "exe_123",
                "matrix_label": "M",
                "context_target": "C",
                "distribution": "dist",
                "hits": 1,
                "total_atoms": 1,
                "hit_ratio": 1.0,
                "raw_score": 1.0,
                "scale_max": 5.0,
                "normalized_score": 20.0,
                "level": 1,
                "level_name": "L1",
                "criterion": "Crit",
                "claim_type": "Claim",
                "result_status": 1,
                "quotes": "Quotes",
                "ai_reasoning": "Reasoning",
                "unknown_extra_field": "disallowed",
            }
        )


@pytest.mark.asyncio
async def test_export_flat_csv_fails_empty_results() -> None:
    """Boundary test: Exporting flat CSV with empty results raises AppException."""
    service = ExportService()
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED, has_atoms=False)
    report_dto = ReportDataDTO(
        workflow_id="wor_0123456789abcdef",
        execution_id="exe_0123456789abcdef",
        profile_id="prf_0123456789abcdef",
        results=[],
    )
    with pytest.raises(AppException) as exc_info:
        await service.export_flat_csv(execution=exec_record, report_dto=report_dto, matrices=[])
    assert exc_info.value.status_code == 400
    assert "Execution has no scoreable atoms" in exc_info.value.message


def _build_test_causal_payloads() -> tuple[CausalGraphPayloadDTO, CausalTdaFusionResultDTO]:
    """Helper constructing validated CausalGraphPayloadDTO and CausalTdaFusionResultDTO."""
    node1 = CausalNodeDTO(
        node_id="tda_0000000000000001",
        claim="Foundational architecture assumption holds true.",
        source_quotes="Foundational verified quote.",
        paragraph_ref="Section 1",
        status=ExecutionStatus.PASSED,
        blame_parent_ids=[],
        dependent_child_ids=["tda_0000000000000002"],
        tda_score=5.0,
        tda_level_key="L5",
    )
    node2 = CausalNodeDTO(
        node_id="tda_0000000000000002",
        claim="Downstream interface verification.",
        source_quotes="Interface failure quote.",
        paragraph_ref="Section 2",
        status=ExecutionStatus.FAILED,
        blame_parent_ids=["tda_0000000000000001"],
        dependent_child_ids=[],
        tda_score=1.0,
        tda_level_key="L1",
    )
    edge = CausalEdgeDTO(
        source_id="tda_0000000000000001",
        target_id="tda_0000000000000002",
        reasoning="Dependency implies interface requirement.",
    )
    graph = CausalGraphPayloadDTO(
        nodes=[node1, node2],
        edges=[edge],
        root_cause_node_ids=["tda_0000000000000001"],
        cycle_detected=False,
        isolated_node_ids=[],
    )
    diag = CausalRootCauseDiagnosisDTO(
        failing_node_id="tda_0000000000000002",
        root_cause_parent_id="tda_0000000000000001",
        paragraph_ref="Section 2",
        explanation="Root assumption lacked schema boundary.",
        penalty_points=1.5,
    )
    audit = AntiFluffAuditDTO(
        isolated_concept_count=0,
        cycle_detected=False,
        unsupported_claims_count=1,
        logical_cohesion_score=0.9,
        detected_hollow_terms=[],
    )
    remediation = PrescriptiveRemediationDTO(
        target_node_id="tda_0000000000000002",
        required_action="Add extra='forbid' to model.",
        unlocks_node_ids=["tda_0000000000000001"],
        potential_score_impact=1.5,
    )
    fair_scoring = FairScoringBreakdownDTO(
        independent_error_count=1,
        cascading_fault_count=1,
        raw_penalty_points=3.0,
        deduplicated_penalty_points=1.5,
        dampened_savings_points=1.5,
    )
    fusion_result = CausalTdaFusionResultDTO(
        graph=graph,
        root_cause_diagnoses=[diag],
        anti_fluff_audit=audit,
        prescriptive_remediations=[remediation],
        fair_scoring=fair_scoring,
    )
    return graph, fusion_result


def test_extract_causal_result_none_and_empty() -> None:
    """Verify _extract_causal_result returns None for None or trace with no causal data."""
    assert ExportService._extract_causal_result(None) is None

    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED, has_atoms=True)
    assert ExportService._extract_causal_result(exec_record) is None


def test_extract_causal_result_from_trace_dict_fusion_and_graph() -> None:
    """Verify _extract_causal_result parses dict with 'fusion', 'graph', or direct 'nodes'."""
    graph, fusion = _build_test_causal_payloads()

    # 1. content is dict containing "fusion"
    exec_fusion = ExecutionRecord(
        id="exe_0123456789abcdef",
        workflow_id="wor_0123456789abcdef",
        organization_id="org_0123456789abcdef",
        target_locale="en",
        status=ExecutionStatus.PASSED,
        execution_trace=[
            TraceEvent(
                step_name="causal_discovery",
                event_type="output",
                content={"fusion": fusion.model_dump(mode="json")},
            )
        ],
    )
    res_fusion = ExportService._extract_causal_result(exec_fusion)
    assert isinstance(res_fusion, CausalTdaFusionResultDTO)
    assert len(res_fusion.graph.nodes) == 2

    # 2. content is dict containing "graph"
    exec_graph = ExecutionRecord(
        id="exe_0123456789abcdef",
        workflow_id="wor_0123456789abcdef",
        organization_id="org_0123456789abcdef",
        target_locale="en",
        status=ExecutionStatus.PASSED,
        execution_trace=[
            TraceEvent(
                step_name="causal_discovery",
                event_type="output",
                content={"graph": graph.model_dump(mode="json")},
            )
        ],
    )
    res_graph = ExportService._extract_causal_result(exec_graph)
    assert isinstance(res_graph, CausalGraphPayloadDTO)
    assert len(res_graph.nodes) == 2

    # 3. content is dict directly containing "nodes"
    exec_raw_graph = ExecutionRecord(
        id="exe_0123456789abcdef",
        workflow_id="wor_0123456789abcdef",
        organization_id="org_0123456789abcdef",
        target_locale="en",
        status=ExecutionStatus.PASSED,
        execution_trace=[
            TraceEvent(
                step_name="causal_discovery",
                event_type="output",
                content=graph.model_dump(mode="json"),
            )
        ],
    )
    res_raw = ExportService._extract_causal_result(exec_raw_graph)
    assert isinstance(res_raw, CausalGraphPayloadDTO)
    assert len(res_raw.nodes) == 2


def test_extract_causal_result_from_step_states() -> None:
    """Verify _extract_causal_result extracts from step_states when trace does not contain it."""
    graph, fusion = _build_test_causal_payloads()

    # 1. StepOutputDTO with CausalTdaFusionResultDTO payload
    exec_step_fusion = ExecutionRecord.model_construct(
        id="exe_0123456789abcdef",
        workflow_id="wor_0123456789abcdef",
        organization_id="org_0123456789abcdef",
        target_locale="en",
        status=ExecutionStatus.PASSED,
        step_states={
            "causal_step": StepOutputDTO(
                step_id="stp_0123456789abcdef",
                block_id="blk_0123456789abcdef",
                data_type="causal",
                payload=fusion,
            )
        },
    )
    assert isinstance(ExportService._extract_causal_result(exec_step_fusion), CausalTdaFusionResultDTO)

    # 2. StepOutputDTO with CausalGraphPayloadDTO payload
    exec_step_graph = ExecutionRecord.model_construct(
        id="exe_0123456789abcdef",
        workflow_id="wor_0123456789abcdef",
        organization_id="org_0123456789abcdef",
        target_locale="en",
        status=ExecutionStatus.PASSED,
        step_states={
            "causal_step": StepOutputDTO(
                step_id="stp_0123456789abcdef",
                block_id="blk_0123456789abcdef",
                data_type="causal",
                payload=graph,
            )
        },
    )
    assert isinstance(ExportService._extract_causal_result(exec_step_graph), CausalGraphPayloadDTO)

    # 3. Direct causal DTO in step_states
    exec_direct = ExecutionRecord.model_construct(
        id="exe_0123456789abcdef",
        workflow_id="wor_0123456789abcdef",
        organization_id="org_0123456789abcdef",
        target_locale="en",
        status=ExecutionStatus.PASSED,
        step_states={"causal_step": graph},
    )
    assert isinstance(ExportService._extract_causal_result(exec_direct), CausalGraphPayloadDTO)

    # 4. StepOutputDTO with non-causal payload
    exec_other = ExecutionRecord.model_construct(
        id="exe_0123456789abcdef",
        workflow_id="wor_0123456789abcdef",
        organization_id="org_0123456789abcdef",
        target_locale="en",
        status=ExecutionStatus.PASSED,
        step_states={
            "text_step": StepOutputDTO(
                step_id="stp_0123456789abcdef",
                block_id="blk_0123456789abcdef",
                data_type="text",
                payload="unrelated text",
            )
        },
    )
    assert ExportService._extract_causal_result(exec_other) is None


def test_resolve_causal_meta_all_branches() -> None:
    """Verify _resolve_causal_meta covers orphan, validated, cascading_fault, root_cause, and None."""
    # 1. c_node is None
    meta_none = ExportService._resolve_causal_meta(None, set(), set())
    assert meta_none.causal_status is None
    assert meta_none.blame_parent_id is None
    assert meta_none.dependent_count is None

    # 2. Isolated set -> orphan
    node_orphan = CausalNodeDTO(
        node_id="tda_0000000000000003",
        claim="Claim orphan",
        source_quotes="Quote",
        paragraph_ref="Section 1",
        status=ExecutionStatus.PASSED,
        blame_parent_ids=[],
        dependent_child_ids=[],
        tda_score=3.0,
        tda_level_key="L3",
    )
    meta_orphan = ExportService._resolve_causal_meta(node_orphan, set(), {"tda_0000000000000003"})
    assert meta_orphan.causal_status == "orphan"

    # 3. Passed -> validated
    node_passed = CausalNodeDTO(
        node_id="tda_0000000000000001",
        claim="Claim passed",
        source_quotes="Quote",
        paragraph_ref="Section 1",
        status=ExecutionStatus.PASSED,
        blame_parent_ids=[],
        dependent_child_ids=["tda_0000000000000002"],
        tda_score=5.0,
        tda_level_key="L5",
    )
    meta_passed = ExportService._resolve_causal_meta(node_passed, set(), set())
    assert meta_passed.causal_status == "validated"
    assert meta_passed.dependent_count == 1

    # 4. Failed with blame parents -> cascading_fault
    node_cascading = CausalNodeDTO(
        node_id="tda_0000000000000002",
        claim="Claim cascading",
        source_quotes="Quote",
        paragraph_ref="Section 2",
        status=ExecutionStatus.FAILED,
        blame_parent_ids=["tda_0000000000000001"],
        dependent_child_ids=[],
        tda_score=1.0,
        tda_level_key="L1",
    )
    meta_cascading = ExportService._resolve_causal_meta(node_cascading, set(), set())
    assert meta_cascading.causal_status == "cascading_fault"
    assert meta_cascading.blame_parent_id == "tda_0000000000000001"

    # 5. Failed without blame parents -> root_cause
    node_root = CausalNodeDTO(
        node_id="tda_0000000000000004",
        claim="Claim root cause",
        source_quotes="Quote",
        paragraph_ref="Section 3",
        status=ExecutionStatus.FAILED,
        blame_parent_ids=[],
        dependent_child_ids=["tda_0000000000000005"],
        tda_score=1.0,
        tda_level_key="L1",
    )
    meta_root = ExportService._resolve_causal_meta(node_root, set(), set())
    assert meta_root.causal_status == "root_cause"
    assert meta_root.blame_parent_id is None
    assert meta_root.dependent_count == 1


def test_resolve_context_target_and_matrix_label_branches() -> None:
    """Verify _resolve_context_target and _resolve_matrix_label fallback paths."""
    # Context target None parent
    assert ExportService._resolve_context_target(None, "en") == ""

    # Context target with fallback string
    row = MatrixScorecardRowDTO(
        block_id="blk_0123456789abcdef",
        name="Axis 1",
        label_i18n=I18nText(translations={"en": "Axis 1"}),
        score=3.0,
        scale_max=5.0,
        row_explanation="Explanation",
        is_evaluative=True,
        context_target="Explicit Fallback Target",
    )
    assert ExportService._resolve_context_target(row, "en") == "Explicit Fallback Target"

    # _resolve_matrix_label parent_m name fallback when label_i18n resolved
    row_label = MatrixScorecardRowDTO(
        block_id="blk_0123456789abcdef",
        name="Fallback Matrix Name",
        label_i18n=I18nText(translations={"en": "Parent Matrix Label", "fi": "Suomeksi"}),
        score=3.0,
        scale_max=5.0,
        row_explanation="Explanation",
        is_evaluative=True,
    )
    resolved_name = ExportService._resolve_matrix_label(
        tda_id="tda_0123456789abcdef",
        atom_matrix_id="blk_unregistered_id",
        parent_m=row_label,
        matrix_title_lookup={},
        blocks_by_id={},
        atom_scale_meta={},
        locale="en",
    )
    assert resolved_name == "Parent Matrix Label"

    # _resolve_matrix_label atom_scale_meta path when atom_matrix_id and parent_m are None
    scale_meta = {
        "tda_0123456789abcdef": AtomScaleMetadataDTO(
            matrix_id="blk_0123456789abcdef",
            level=1,
            level_name="Level 1",
            criterion="Criterion",
            is_inverse=False,
        )
    }
    resolved_scale = ExportService._resolve_matrix_label(
        tda_id="tda_0123456789abcdef",
        atom_matrix_id=None,
        parent_m=None,
        matrix_title_lookup={"blk_0123456789abcdef": "Lookup Title"},
        blocks_by_id={},
        atom_scale_meta=scale_meta,
        locale="en",
    )
    assert resolved_scale == "Lookup Title"


def test_build_matrix_summary_rows_extra_branches() -> None:
    """Verify _build_summary_dtos handles raw_score, cited_source_id, and localized label."""
    service = ExportService()
    row = MatrixScorecardRowDTO(
        block_id="blk_0123456789abcdef",
        name="Fallback Axis Name",
        label_i18n=I18nText(translations={"en": "English Axis", "fi": "Akseli"}),
        score=4.0,
        raw_score=3.5,
        scale_max=5.0,
        row_explanation="Summary reasoning",
        is_evaluative=True,
        cited_source_id="[src_doc_1]",
    )
    rows = service._build_summary_dtos([row], "en", "exe_0123456789abcdef")
    assert len(rows) == 1
    assert rows[0].label == "English Axis"
    assert rows[0].raw_score == 3.5
    assert rows[0].source == "src_doc_1"


@pytest.mark.asyncio
async def test_export_excel_includes_causal_propositions_with_metrics() -> None:
    """Verify Excel Tab 2 includes unmapped causal propositions with causal metrics and ML-ready columns."""
    import io

    import openpyxl

    from backend_v2.models.dtos.atom_result import AtomResultDTO, HydratedAtomDTO
    from backend_v2.models.dtos.causal_discovery import CausalGraphPayloadDTO, CausalNodeDTO
    from backend_v2.models.enums import ExecutionStatus

    service = ExportService()
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED)

    # 1 mapped atom
    atom_res = AtomResultDTO(
        tda_id="tda_0000000000000001",
        matrix_id="blk_0123456789abcdef",
        status=ExecutionStatus.PASSED,
        source_quotes="Quote 1",
        evaluation_reasoning="Affirmative claim passed",
        is_inverse_evidence=False,
    )
    hydrated = HydratedAtomDTO(
        sdui_component=LaxSDUIComponentType.BOOLEAN_CARD,
        resolved_claim="Atom 1 Claim",
    )
    causal_payload = CausalGraphPayloadDTO(
        nodes=[
            CausalNodeDTO(
                node_id="tda_0000000000000001",
                claim="Atom 1 Claim",
                paragraph_ref="p1",
                status=ExecutionStatus.PASSED,
                blame_parent_ids=[],
                dependent_child_ids=["causal_prop_0001"],
            ),
            CausalNodeDTO(
                node_id="causal_prop_0001",
                claim="Unmapped Causal Proposition",
                paragraph_ref="p2",
                status=ExecutionStatus.FAILED,
                blame_parent_ids=["tda_0000000000000001"],
                dependent_child_ids=[],
            ),
        ],
        edges=[],
        root_cause_node_ids=[],
        cycle_detected=False,
        isolated_node_ids=[],
    )
    exec_record = exec_record.model_copy(
        update={
            "execution_trace": [
                TraceEvent(
                    step_name="causal_discovery",
                    event_type="output",
                    content=causal_payload,
                )
            ]
        }
    )
    report_dto = _build_sample_report_dto()
    report_dto = report_dto.model_copy(
        update={
            "results": [atom_res],
            "hydrated_references": {"tda_0000000000000001": hydrated},
        }
    )
    matrices = [
        MatrixScorecardRowDTO(
            block_id="blk_0123456789abcdef",
            name="Matrix 1",
            label_i18n=I18nText(translations={"fi": "Matriisi 1", "en": "Matrix 1"}),
            score=5.0,
            scale_max=5.0,
            row_explanation="Explanation",
            is_evaluative=True,
            true_atoms=1,
            total_atoms=1,
        )
    ]
    payload = await service.export_excel(
        execution=exec_record,
        report_dto=report_dto,
        matrices=matrices,
        locale="fi",
    )
    wb = openpyxl.load_workbook(io.BytesIO(payload.content_bytes))
    tab2 = wb["Raakadata"]
    rows = list(tab2.iter_rows(values_only=True))
    assert len(rows[0]) == 14
    # Header check
    assert rows[0][0] == "Logiikkamatriisi"
    assert rows[0][6] == "Käänteinen ehto"
    assert rows[0][7] == "Saadut pisteet"
    assert rows[0][11] == "Kausaalistatus"
    assert rows[0][12] == "Juurisyyn tunnus"
    assert rows[0][13] == "Riippuvuuksien lkm"

    # Row 1 (mapped atom)
    assert rows[1][0] == "Matriisi 1"
    assert rows[1][6] == 0
    assert rows[1][7] == 1.0
    assert rows[1][11] == "validated"
    assert rows[1][13] == 1

    # Row 2 (unmapped causal proposition)
    assert rows[2][0] == "Kausaalipäättely"
    assert rows[2][2] == 1  # level
    assert rows[2][3] == "Propositio"  # level_name
    assert rows[2][4] == "Unmapped Causal Proposition"  # criterion
    assert rows[2][6] == 0  # is_inverse
    assert rows[2][7] == 0.0  # points_awarded
    assert rows[2][11] == "cascading_fault"  # causal_status
    assert rows[2][12] == "tda_0000000000000001"  # blame_parent_id
    assert rows[2][13] == 0  # dependent_count


@pytest.mark.asyncio
async def test_export_csv_includes_causal_and_ml_columns() -> None:
    """Verify CSV export contains exactly 22 columns with ML indicators and unmapped causal propositions."""
    import io

    from backend_v2.models.dtos.atom_result import AtomResultDTO, HydratedAtomDTO
    from backend_v2.models.dtos.causal_discovery import CausalGraphPayloadDTO, CausalNodeDTO
    from backend_v2.models.enums import ExecutionStatus

    service = ExportService()
    exec_record = _build_sample_execution(status=ExecutionStatus.PASSED)

    atom_res_1 = AtomResultDTO(
        tda_id="tda_0000000000000001",
        matrix_id="blk_0123456789abcdef",
        status=ExecutionStatus.PASSED,
        source_quotes="Quote 1",
        evaluation_reasoning="Affirmative passed",
        is_inverse_evidence=False,
    )
    hydrated_1 = HydratedAtomDTO(
        sdui_component=LaxSDUIComponentType.BOOLEAN_CARD,
        resolved_claim="Criterion 1",
    )
    atom_res_2 = AtomResultDTO(
        tda_id="tda_0000000000000002",
        matrix_id="blk_0123456789abcdef",
        status=ExecutionStatus.PASSED,
        source_quotes=None,
        evaluation_reasoning="Inverse claim passed (anti-pattern absent = clean finding)",
        is_inverse_evidence=True,
    )
    hydrated_2 = HydratedAtomDTO(
        sdui_component=LaxSDUIComponentType.BOOLEAN_CARD,
        resolved_claim="Anti-pattern 1",
    )

    causal_payload_csv = CausalGraphPayloadDTO(
        nodes=[
            CausalNodeDTO(
                node_id="causal_prop_0002",
                claim="Standalone Causal Node",
                paragraph_ref="p3",
                status=ExecutionStatus.SYSTEM_ERROR,
                blame_parent_ids=[],
                dependent_child_ids=[],
            ),
        ],
        edges=[],
        root_cause_node_ids=["causal_prop_0002"],
        cycle_detected=False,
        isolated_node_ids=[],
    )
    exec_record = exec_record.model_copy(
        update={
            "execution_trace": [
                TraceEvent(
                    step_name="causal_discovery",
                    event_type="output",
                    content=causal_payload_csv,
                )
            ]
        }
    )
    report_dto = _build_sample_report_dto()
    report_dto = report_dto.model_copy(
        update={
            "results": [atom_res_1, atom_res_2],
            "hydrated_references": {
                "tda_0000000000000001": hydrated_1,
                "tda_0000000000000002": hydrated_2,
            },
        }
    )
    matrices = [
        MatrixScorecardRowDTO(
            block_id="blk_0123456789abcdef",
            name="Matrix 1",
            label_i18n=I18nText(translations={"en": "Matrix 1"}),
            score=5.0,
            scale_max=5.0,
            row_explanation="Explanation",
            is_evaluative=True,
            true_atoms=1,
            total_atoms=2,
        )
    ]
    payload = await service.export_flat_csv(
        execution=exec_record,
        report_dto=report_dto,
        matrices=matrices,
        locale="en",
    )
    csv_text = payload.content_bytes.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(csv_text)))
    assert len(rows[0]) == 22
    assert rows[0] == ReportHeaderResolver.get_all_denormalized_csv_headers("en")

    # Row 1 (affirmative atom)
    assert rows[1][14] == "0"  # is_inverse
    assert float(rows[1][15]) == 1.0  # points_awarded

    # Row 2 (inverse atom: anti-pattern passed means clean finding, points_awarded = 0.0)
    assert rows[2][14] == "1"  # is_inverse
    assert float(rows[2][15]) == 0.0  # points_awarded

    # Mathematical scoring invariant: sum(points_awarded) for affirmative == hits
    affirmative_points = sum(float(r[15]) for r in rows[1:3] if r[14] == "0")
    assert affirmative_points == matrices[0].true_atoms

    # Row 3 (unmapped causal proposition)
    assert rows[3][1] == "Causal Inference"
    assert rows[3][11] == "Proposition"
    assert rows[3][12] == "Standalone Causal Node"
    assert rows[3][14] == "0"  # is_inverse
    assert float(rows[3][15]) == 0.0  # points_awarded
    assert rows[3][19] == "root_cause"  # causal_status
