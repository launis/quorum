from __future__ import annotations

import csv
import io
import json
from pathlib import Path

import openpyxl
import pytest
from pydantic import ValidationError

from backend_v2.models.core_base import I18nText
from backend_v2.models.domain.execution import ExecutionRecord, ExecutionStepState
from backend_v2.models.dtos.atom_result import AtomResultDTO, HydratedAtomDTO
from backend_v2.models.dtos.export import (
    ExportForensicAtomDTO,
    ExportMatrixSummaryRowDTO,
    ExportPayloadDTO,
)
from backend_v2.models.dtos.matrix_scorecard import MatrixScorecardRowDTO
from backend_v2.models.dtos.report_data import ReportDataDTO
from backend_v2.models.enums import (
    ExecutionStatus,
    ReportAtomColumn,
    ReportMatrixColumn,
    ReportSheetKey,
)
from backend_v2.services.export_service import ExportService
from backend_v2.services.localization import LocalizationService, ReportHeaderResolver


def _load_json_l10n(locale: str) -> dict[str, str]:
    """Load backend JSON localization dictionary."""
    path = Path("backend_v2/l10n") / f"{locale}.json"
    with open(path, encoding="utf-8") as f:
        data: dict[str, str] = json.load(f)
    return data


def _load_flutter_arb(locale: str) -> dict[str, str]:
    """Load Flutter ARB localization dictionary."""
    path = Path("client_app_v2/lib/l10n") / f"app_{locale}.arb"
    with open(path, encoding="utf-8") as f:
        raw = json.load(f)
        data: dict[str, str] = {str(k): str(v) for k, v in raw.items()}
    return data


def test_report_matrix_columns_translation_completeness() -> None:
    """Assert every member of ReportMatrixColumn has non-empty translations in fi and en."""
    for locale in ["fi", "en"]:
        l10n = _load_json_l10n(locale)
        for col in ReportMatrixColumn:
            assert col.l10n_key in l10n, f"Missing key {col.l10n_key} in {locale}.json"
            assert l10n[col.l10n_key].strip() != "", f"Empty key {col.l10n_key} in {locale}.json"
            header = ReportHeaderResolver.get_matrix_column_header(col, locale=locale)
            assert header == l10n[col.l10n_key]


def test_report_atom_columns_translation_completeness() -> None:
    """Assert every member of ReportAtomColumn has non-empty translations in fi and en."""
    for locale in ["fi", "en"]:
        l10n = _load_json_l10n(locale)
        for col in ReportAtomColumn:
            assert col.l10n_key in l10n, f"Missing key {col.l10n_key} in {locale}.json"
            assert l10n[col.l10n_key].strip() != "", f"Empty key {col.l10n_key} in {locale}.json"
            header = ReportHeaderResolver.get_atom_column_header(col, locale=locale)
            assert header == l10n[col.l10n_key]


def test_report_sheet_keys_translation_completeness() -> None:
    """Assert every member of ReportSheetKey has non-empty translations in fi and en."""
    for locale in ["fi", "en"]:
        l10n = _load_json_l10n(locale)
        for sheet in ReportSheetKey:
            assert sheet.l10n_key in l10n, f"Missing key {sheet.l10n_key} in {locale}.json"
            assert l10n[sheet.l10n_key].strip() != "", f"Empty key {sheet.l10n_key} in {locale}.json"
            name = ReportHeaderResolver.get_sheet_name(sheet, locale=locale)
            assert name == l10n[sheet.l10n_key]


def test_cross_platform_studio_ui_arb_parity() -> None:
    """Assert 1:1 cross-platform parity between Flutter Studio ARB files and Backend JSON."""
    matrix_col_mappings: list[tuple[str, str]] = [
        ("studioMatrixColLabel", "matrix_col_label"),
        ("studioMatrixColContextTarget", "matrix_col_context_target"),
        ("studioMatrixColDistribution", "matrix_col_distribution"),
        ("studioMatrixColRowExplanation", "matrix_col_row_explanation"),
        ("studioMatrixColCriteria", "matrix_col_criteria"),
        ("studioMatrixColQuotes", "matrix_col_quotes"),
        ("studioMatrixColSource", "matrix_col_source"),
        ("studioMatrixColNormalized", "matrix_col_normalized_score"),
        ("studioMatrixColScore", "matrix_col_score"),
    ]

    metadata_mappings: list[tuple[str, str]] = [
        ("metaDate", "metadata_date"),
        ("metaOrganization", "metadata_organization"),
        ("metaUser", "metadata_user"),
        ("metaScoringEngine", "metadata_scoring_engine"),
        ("metaStrictness", "metadata_strictness"),
        ("metaCost", "metadata_cost"),
        ("metaTokens", "metadata_tokens"),
    ]

    xai_mappings: list[tuple[str, str]] = [
        ("xaiJustification", "xai_ext_justification"),
        ("xaiCoachingTip", "xai_ext_coaching"),
        ("xaiDevilsAdvocate", "xai_ext_falsification"),
        ("xaiMissingContext", "xai_ext_missing_context"),
        ("xaiRiskFlag", "xai_ext_risk_flag"),
        ("xaiRemediation", "xai_ext_remediation_steps"),
        ("xaiSentiment", "xai_ext_emotional_sentiment"),
        ("xaiTheoryLink", "xai_ext_theory_link"),
        ("xaiConfidence", "xai_ext_confidence"),
        ("xaiSourceCitation", "xai_ext_citation"),
        ("xaiContextualOverride", "xai_ext_contextual_override"),
        ("xaiSourceId", "xai_ext_source_id"),
    ]

    for locale in ["fi", "en"]:
        json_l10n = _load_json_l10n(locale)
        arb_l10n = _load_flutter_arb(locale)

        # 1. Assert all 9 matrix columns match 1:1
        for arb_key, json_key in matrix_col_mappings:
            assert arb_key in arb_l10n, f"Missing ARB key {arb_key} in app_{locale}.arb"
            assert json_key in json_l10n, f"Missing JSON key {json_key} in {locale}.json"
            assert arb_l10n[arb_key] == json_l10n[json_key], (
                f"Mismatch for {arb_key} vs {json_key} in {locale}: '{arb_l10n[arb_key]}' != '{json_l10n[json_key]}'"
            )

        # 2. Assert all 7 metadata fields match 1:1
        for arb_key, json_key in metadata_mappings:
            assert arb_key in arb_l10n, f"Missing ARB key {arb_key} in app_{locale}.arb"
            assert json_key in json_l10n, f"Missing JSON key {json_key} in {locale}.json"
            assert arb_l10n[arb_key] == json_l10n[json_key], (
                f"Mismatch for {arb_key} vs {json_key} in {locale}: '{arb_l10n[arb_key]}' != '{json_l10n[json_key]}'"
            )

        # 3. Assert all 12 XAI extension labels match 1:1
        for arb_key, json_key in xai_mappings:
            assert arb_key in arb_l10n, f"Missing ARB key {arb_key} in app_{locale}.arb"
            assert json_key in json_l10n, f"Missing JSON key {json_key} in {locale}.json"
            assert arb_l10n[arb_key] == json_l10n[json_key], (
                f"Mismatch for {arb_key} vs {json_key} in {locale}: '{arb_l10n[arb_key]}' != '{json_l10n[json_key]}'"
            )


def test_cross_surface_header_parity() -> None:
    """Verify that ReportHeaderResolver resolves identical headers for Excel and CSV."""
    for locale in ["fi", "en"]:
        matrix_headers = ReportHeaderResolver.get_all_matrix_headers(locale=locale)
        atom_headers = ReportHeaderResolver.get_all_atom_headers(locale=locale)

        assert len(matrix_headers) == len(ReportMatrixColumn) == 13
        assert len(atom_headers) == len(ReportAtomColumn) == 9

        for col in ReportMatrixColumn:
            assert matrix_headers[col] == ReportHeaderResolver.get_matrix_column_header(col, locale=locale)

        for idx, col in enumerate(ReportAtomColumn):
            assert atom_headers[idx] == ReportHeaderResolver.get_atom_column_header(col, locale=locale)


def test_dto_field_mapping_parity() -> None:
    """Verify ExportMatrixSummaryRowDTO and ExportForensicAtomDTO fields map directly to column enums."""
    summary_fields = set(ExportMatrixSummaryRowDTO.model_fields.keys())
    matrix_enum_values = {col.value for col in ReportMatrixColumn}
    assert summary_fields == matrix_enum_values, (
        f"ExportMatrixSummaryRowDTO fields do not match ReportMatrixColumn: {summary_fields} != {matrix_enum_values}"
    )

    atom_fields = set(ExportForensicAtomDTO.model_fields.keys())
    expected_atom_fields = {
        "matrix_label",
        "context_target",
        "level",
        "level_name",
        "criterion",
        "claim_type",
        "result_status",
        "quotes",
        "ai_reasoning",
    }
    assert atom_fields == expected_atom_fields


def test_pydantic_strict_extra_forbid_gate() -> None:
    """Verify that attempting to instantiate export DTOs with undeclared fields raises ValidationError."""
    with pytest.raises(ValidationError):
        ExportForensicAtomDTO.model_validate(
            {
                "matrix_label": "M",
                "context_target": "C",
                "level": 1,
                "level_name": "L1",
                "criterion": "Crit",
                "claim_type": "Claim",
                "result_status": 1,
                "quotes": "Q",
                "ai_reasoning": "Reason",
                "synthetic_extra_field": "malicious",
            }
        )

    with pytest.raises(ValidationError):
        ExportMatrixSummaryRowDTO.model_validate(
            {
                "execution_id": "exe_123",
                "label": "L",
                "context_target": "C",
                "distribution": "D",
                "hits": 1,
                "total_atoms": 1,
                "hit_ratio": 1.0,
                "raw_score": 1.0,
                "scale_max": 5.0,
                "normalized_score": 20.0,
                "row_explanation": "E",
                "criteria": "1/1",
                "source": "S",
                "phantom_column": "forbidden",
            }
        )

    with pytest.raises(ValidationError):
        ExportPayloadDTO.model_validate(
            {
                "content_bytes": b"data",
                "filename": "export.xlsx",
                "mime_type": "application/vnd.ms-excel",
                "extra_token": "rejected",
            }
        )


def test_unsupported_locale_fallback() -> None:
    """Verify that passing an unsupported locale code deterministically falls back to en."""
    header_de = ReportHeaderResolver.get_matrix_column_header(ReportMatrixColumn.LABEL, locale="de")
    header_en = ReportHeaderResolver.get_matrix_column_header(ReportMatrixColumn.LABEL, locale="en")
    assert header_de == header_en

    atom_header_fr = ReportHeaderResolver.get_atom_column_header(ReportAtomColumn.MATRIX, locale="fr")
    atom_header_en = ReportHeaderResolver.get_atom_column_header(ReportAtomColumn.MATRIX, locale="en")
    assert atom_header_fr == atom_header_en

    sheet_ja = ReportHeaderResolver.get_sheet_name(ReportSheetKey.SUMMARY, locale="ja")
    sheet_en = ReportHeaderResolver.get_sheet_name(ReportSheetKey.SUMMARY, locale="en")
    assert sheet_ja == sheet_en


@pytest.mark.asyncio
async def test_mathematical_dumb_painter_invariance() -> None:
    """Verify exact Dumb Painter row and column count invariance across Excel Tab 1, Tab 2, and CSV."""
    service = ExportService()
    execution = ExecutionRecord(
        id="exe_0123456789abcdef",
        workflow_id="wor_0123456789abcdef",
        organization_id="org_0123456789abcdef",
        target_locale="fi",
        status=ExecutionStatus.PASSED,
        step_states={
            "stp_1": ExecutionStepState(
                id="stp_1",
                label="Step 1",
                status=ExecutionStatus.PASSED,
            )
        },
    )

    matrices = [
        MatrixScorecardRowDTO(
            block_id="blk_0123456789abcdef",
            name="Logiikkamatriisi Alpha",
            label_i18n=I18nText(translations={"fi": "Logiikkamatriisi Alpha", "en": "Logic Matrix Alpha"}),
            score=8.5,
            scale_max=10.0,
            normalized_score=85.0,
            row_explanation="Selitys 1",
            is_evaluative=True,
        ),
        MatrixScorecardRowDTO(
            block_id="blk_fedcba9876543210",
            name="Logiikkamatriisi Beta",
            label_i18n=I18nText(translations={"fi": "Logiikkamatriisi Beta", "en": "Logic Matrix Beta"}),
            score=7.0,
            scale_max=10.0,
            normalized_score=70.0,
            row_explanation="Selitys 2",
            is_evaluative=True,
        ),
    ]

    report_results = [
        AtomResultDTO(
            tda_id="tda_0000000000000001",
            matrix_id="blk_0123456789abcdef",
            status=ExecutionStatus.PASSED,
            source_quote="Lainaus 1",
            evaluation_reasoning="Pätevä perustelu 1",
        ),
        AtomResultDTO(
            tda_id="tda_0000000000000002",
            matrix_id="blk_0123456789abcdef",
            status=ExecutionStatus.FAILED,
            source_quote=None,
            evaluation_reasoning="Hylätty perustelu 2",
        ),
        AtomResultDTO(
            tda_id="tda_0000000000000003",
            matrix_id="blk_fedcba9876543210",
            status=ExecutionStatus.PASSED,
            source_quote="Lainaus 3",
            evaluation_reasoning="Pätevä perustelu 3",
        ),
    ]

    hydrated_refs = {
        "tda_0000000000000001": HydratedAtomDTO(
            sdui_component="boolean_card",
            resolved_claim="Väite 1",
        ),
        "tda_0000000000000002": HydratedAtomDTO(
            sdui_component="boolean_card",
            resolved_claim="Väite 2",
        ),
        "tda_0000000000000003": HydratedAtomDTO(
            sdui_component="boolean_card",
            resolved_claim="Väite 3",
        ),
    }

    report_dto = ReportDataDTO(
        workflow_id="wor_0123456789abcdef",
        execution_id="exe_0123456789abcdef",
        profile_id="prf_0123456789abcdef",
        results=report_results,
        hydrated_references=hydrated_refs,
    )

    # 1. Test Excel Invariance
    excel_payload = await service.export_excel(
        execution=execution,
        report_dto=report_dto,
        matrices=matrices,
        locale="fi",
    )
    assert isinstance(excel_payload, ExportPayloadDTO)
    wb = openpyxl.load_workbook(io.BytesIO(excel_payload.content_bytes))

    sheet1 = wb[LocalizationService.translate("export_sheet_summary", "fi")]
    sheet1_rows = list(sheet1.iter_rows(values_only=True))
    assert len(sheet1_rows) == len(matrices) + 1  # 1 header row + N data rows
    assert len(sheet1_rows[0]) == len(ReportMatrixColumn) == 13

    sheet2 = wb[LocalizationService.translate("export_sheet_raw_data", "fi")]
    sheet2_rows = list(sheet2.iter_rows(values_only=True))
    assert len(sheet2_rows) == len(report_dto.results) + 1  # 1 header row + M data rows
    assert len(sheet2_rows[0]) == len(ReportAtomColumn) == 9

    # Verify binary integer status on Tab 2
    status_col_idx = 6  # ReportAtomColumn.RESULT_STATUS index
    for row in sheet2_rows[1:]:
        assert row[status_col_idx] in (0, 1), f"Status {row[status_col_idx]} is not strictly 0 or 1"

    # 2. Test CSV Invariance
    csv_payload = await service.export_flat_csv(
        execution=execution,
        report_dto=report_dto,
        matrices=matrices,
        locale="fi",
    )
    assert isinstance(csv_payload, ExportPayloadDTO)
    csv_text = csv_payload.content_bytes.decode("utf-8-sig")
    reader = list(csv.reader(io.StringIO(csv_text)))

    assert len(reader) == len(report_dto.results) + 1  # 1 header row + M data rows
    assert len(reader[0]) == 17

    # Verify binary integer status on CSV
    csv_status_col_idx = 14
    for row in reader[1:]:
        assert row[csv_status_col_idx] in ("0", "1"), f"CSV status {row[csv_status_col_idx]} is not strictly 0 or 1"


def test_tabular_rows_tab_arb_and_backend_parity() -> None:
    """Assert 1:1 cross-platform parity between Flutter UI table column keys and Backend SSOT keys."""
    tabular_mappings: list[tuple[str, str]] = [
        ("tableColumnCriteriaMetric", "matrix_col_criteria"),
        ("tableColumnScore", "matrix_col_score"),
        ("tableColumnReasoningQuote", "matrix_col_row_explanation"),
    ]
    for locale in ["fi", "en"]:
        json_l10n = _load_json_l10n(locale)
        arb_l10n = _load_flutter_arb(locale)
        for arb_key, json_key in tabular_mappings:
            assert arb_key in arb_l10n, f"Missing ARB key {arb_key} in app_{locale}.arb"
            assert json_key in json_l10n, f"Missing JSON key {json_key} in {locale}.json"
            assert arb_l10n[arb_key] == json_l10n[json_key], (
                f"Mismatch for {arb_key} vs {json_key} in {locale}: '{arb_l10n[arb_key]}' != '{json_l10n[json_key]}'"
            )
