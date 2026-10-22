"""Excel styling routines for corporate presentation sheets.

Provides openpyxl styling utilities including header fills, zebra striping,
borders, alignment, and dynamic column auto-width calculation.
"""

from __future__ import annotations

from typing import TYPE_CHECKING

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    import openpyxl

__all__ = ["apply_corporate_excel_styling"]


def apply_corporate_excel_styling(workbook: openpyxl.Workbook, sheet_names: list[str]) -> None:
    """Applies professional corporate styling across export sheets.

    Args:
        workbook: Active openpyxl Workbook instance.
        sheet_names: List of sheet names to style.
    """
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    thin_border_side = Side(border_style="thin", color="D9D9D9")
    grid_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    data_alignment = Alignment(vertical="top", wrap_text=True)

    for name in sheet_names:
        if name not in workbook.sheetnames:
            continue
        ws = workbook[name]
        if ws.views.sheetView:
            ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "A2"

        ws.row_dimensions[1].height = 28
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = grid_border

        for row_idx in range(2, ws.max_row + 1):
            fill = white_fill
            if row_idx % 2 == 0:
                fill = zebra_fill
            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.fill = fill
                cell.border = grid_border
                cell.alignment = data_alignment

        for col_idx in range(1, ws.max_column + 1):
            col_letter = get_column_letter(col_idx)
            max_len = 0
            for row_idx in range(1, ws.max_row + 1):
                val = ws.cell(row=row_idx, column=col_idx).value
                if val is not None:
                    val_str = str(val)
                    line_len = max((len(line) for line in val_str.split("\n")), default=0)
                    max_len = max(max_len, line_len)
            calculated_width = max(max_len + 4, 12)
            ws.column_dimensions[col_letter].width = min(calculated_width, 50)
